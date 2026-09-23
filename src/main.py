from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from ortools.sat.python import cp_model


BASE_DIR = Path(__file__).resolve().parent.parent
INPUT_FILE = BASE_DIR / "input" / "kinmu_sample.xlsx"
OUTPUT_DIR = BASE_DIR / "output"
OUTPUT_FILE = OUTPUT_DIR / "kinmu_output.xlsx"
TEMP_FILE = OUTPUT_DIR / ".kinmu_output.tmp.xlsx"

SHEET_NAME = "Sheet1"

DATE_ROW = 2
WEEKDAY_ROW = 3
DAY_START_COLUMN = 3       # C
DAY_END_COLUMN = 33        # AG

WARD_MEMBERS = {
    1: {"A", "B", "C", "D"},
    2: {"E", "F", "G", "H", "I", "J", "K"},
    3: {"L", "M", "N", "O", "P"},
}

EMPLOYEE_NAMES = list("ABCDEFGHIJKLMNOP")

WORK_VALUES = {
    None,
    "休",
    "出",
    "半/",
    "/半",
}

REASON_VALUES = {
    None,
    "希望",
    "委員会",
    "年休",
    "年/",
    "/年",
    "研修",
    "特休",
    "出張",
}

FIXED_WORK_VALUES = {
    "休",
    "出",
    "半/",
    "/半",
}

FIXED_REASON_VALUES = REASON_VALUES - {None}

WEEKDAYS = {
    "月",
    "火",
    "水",
    "木",
    "金",
    "土",
    "日",
}

MON_TO_SAT = WEEKDAYS - {"日"}


# ソフト条件の優先度
WEIGHT_FOUR_CONSECUTIVE = 500
WEIGHT_LEADER = 400
WEIGHT_OT_MINIMUM = 400
WEIGHT_SUNDAY_STAFF = 400
WEIGHT_TOO_MANY_OFF = 300
WEIGHT_THIRD_OFF = 100
WEIGHT_WARD1 = 100
WEIGHT_WARD2 = 80
WEIGHT_WARD3 = 100
WEIGHT_OT_SECOND = 50
WEIGHT_ISOLATED_WORK = 2


@dataclass(frozen=True)
class Employee:
    name: str
    row: int
    ward: int
    is_dedicated: bool
    is_assigned: bool
    is_ot: bool


@dataclass(frozen=True)
class HolidayTarget:
    public_half_units: int
    paid_half_units: int

    @property
    def total_half_units(self) -> int:
        return self.public_half_units + self.paid_half_units


@dataclass(frozen=True)
class Penalty:
    label: str
    variable: cp_model.IntVar
    weight: int


@dataclass
class ModelData:
    model: cp_model.CpModel
    full_work: dict[tuple[int, int], cp_model.IntVar]
    fixed_shifts: dict[tuple[int, int], str]
    existing_changes: list[tuple[str, cp_model.IntVar]]
    penalties: list[Penalty]
    holiday_targets: dict[str, HolidayTarget]
    holiday_off: dict[str, cp_model.LinearExpr]
    holiday_deviation: dict[str, cp_model.IntVar]
    max_holiday_deviation: cp_model.IntVar


def normalize(value):
    """空文字や前後の空白を None に統一する。"""
    if isinstance(value, str):
        value = value.strip()
        return value or None

    return value


def employee_name(raw_value) -> str | None:
    """A（専従）のような文字列から A だけを取り出す。"""
    value = normalize(raw_value)

    if value is None:
        return None

    return re.split(
        r"[（(]",
        str(value),
        maxsplit=1,
    )[0].strip()


def ward_for(name: str) -> int:
    for ward, members in WARD_MEMBERS.items():
        if name in members:
            return ward

    raise ValueError(
        f"病棟を特定できません: {name}"
    )


def discover_employees(
    sheet: Worksheet,
) -> list[Employee]:
    """A列からA〜Pの職員と役割を取得する。"""
    found: dict[str, Employee] = {}

    for row in range(
        1,
        sheet.max_row + 1,
    ):
        raw = normalize(
            sheet.cell(
                row=row,
                column=1,
            ).value
        )

        name = employee_name(raw)

        if name not in EMPLOYEE_NAMES:
            continue

        if name in found:
            raise ValueError(
                f"職員「{name}」が"
                "A列に複数あります。"
            )

        text = str(raw)

        found[name] = Employee(
            name=name,
            row=row,
            ward=ward_for(name),
            is_dedicated="専従" in text,
            is_assigned="専任" in text,
            is_ot="作業療法士" in text,
        )

    missing = [
        name
        for name in EMPLOYEE_NAMES
        if name not in found
    ]

    if missing:
        raise ValueError(
            "職員が見つかりません: "
            + ", ".join(missing)
        )

    return [
        found[name]
        for name in EMPLOYEE_NAMES
    ]


def day_columns(
    sheet: Worksheet,
) -> list[int]:
    """勤務対象の日付列を取得する。"""
    columns = []

    for column in range(
        DAY_START_COLUMN,
        DAY_END_COLUMN + 1,
    ):
        date_value = normalize(
            sheet.cell(
                row=DATE_ROW,
                column=column,
            ).value
        )

        weekday_value = normalize(
            sheet.cell(
                row=WEEKDAY_ROW,
                column=column,
            ).value
        )

        if (
            date_value is None
            and weekday_value is None
        ):
            continue

        if weekday_value not in WEEKDAYS:
            cell = sheet.cell(
                row=WEEKDAY_ROW,
                column=column,
            ).coordinate

            raise ValueError(
                f"{cell} の曜日を"
                f"認識できません: "
                f"{weekday_value}"
            )

        columns.append(column)

    if not columns:
        raise ValueError(
            "勤務対象の日付列が"
            "見つかりません。"
        )

    return columns


def day_label(
    sheet: Worksheet,
    column: int,
) -> str:
    value = normalize(
        sheet.cell(
            row=DATE_ROW,
            column=column,
        ).value
    )

    if value is not None:
        return str(value)

    return str(
        column
        - DAY_START_COLUMN
        + 1
    )


def weekday(
    sheet: Worksheet,
    column: int,
) -> str:
    value = normalize(
        sheet.cell(
            row=WEEKDAY_ROW,
            column=column,
        ).value
    )

    if value not in WEEKDAYS:
        raise ValueError(
            f"曜日を認識できません: {value}"
        )

    return str(value)


def find_header_column(
    sheet: Worksheet,
    header: str,
) -> int:
    """上部の見出しから列番号を探す。"""
    for row in range(
        1,
        WEEKDAY_ROW + 1,
    ):
        for column in range(
            1,
            sheet.max_column + 1,
        ):
            value = normalize(
                sheet.cell(
                    row=row,
                    column=column,
                ).value
            )

            if value == header:
                return column

    raise ValueError(
        f"見出し「{header}」が"
        "見つかりません。"
    )


def is_fixed_shift(
    work,
    reason,
) -> bool:
    return (
        work in FIXED_WORK_VALUES
        and reason in FIXED_REASON_VALUES
    )


def validate_input(
    sheet: Worksheet,
    employees: list[Employee],
    columns: list[int],
) -> list[str]:
    errors = []

    for employee in employees:
        reason_row = employee.row + 1

        for column in columns:
            work = normalize(
                sheet.cell(
                    employee.row,
                    column,
                ).value
            )

            reason = normalize(
                sheet.cell(
                    reason_row,
                    column,
                ).value
            )

            if work not in WORK_VALUES:
                cell = sheet.cell(
                    employee.row,
                    column,
                ).coordinate

                errors.append(
                    f"{cell}: "
                    f"不明な勤務「{work}」"
                )

            if reason not in REASON_VALUES:
                cell = sheet.cell(
                    reason_row,
                    column,
                ).coordinate

                errors.append(
                    f"{cell}: "
                    f"不明な理由「{reason}」"
                )

            if (
                reason in FIXED_REASON_VALUES
                and work
                not in FIXED_WORK_VALUES
            ):
                cell = sheet.cell(
                    employee.row,
                    column,
                ).coordinate

                errors.append(
                    f"{cell}: "
                    f"理由「{reason}」がありますが"
                    "勤務が未入力です。"
                )

            if (
                reason == "年休"
                and work != "休"
            ):
                cell = sheet.cell(
                    employee.row,
                    column,
                ).coordinate

                errors.append(
                    f"{cell}: "
                    "年休は「休」と"
                    "組み合わせてください。"
                )

            if (
                reason in {"年/", "/年"}
                and work
                not in {"半/", "/半"}
            ):
                cell = sheet.cell(
                    employee.row,
                    column,
                ).coordinate

                errors.append(
                    f"{cell}: "
                    "半日の年休は"
                    "「半/」または「/半」"
                    "と組み合わせてください。"
                )

    return errors


def capture_formulas(
    sheet: Worksheet,
) -> dict[str, str]:
    return {
        cell.coordinate: cell.value
        for row in sheet.iter_rows()
        for cell in row
        if (
            isinstance(cell.value, str)
            and cell.value.startswith("=")
        )
    }


def validate_formulas(
    sheet: Worksheet,
    formulas: dict[str, str],
) -> None:
    for address, expected in formulas.items():
        if sheet[address].value != expected:
            raise RuntimeError(
                f"数式が変更されました: "
                f"{address}"
            )


def validate_fixed_shifts(
    sheet: Worksheet,
    fixed_shifts: dict[
        tuple[int, int],
        str,
    ],
) -> None:
    for (
        row,
        column,
    ), expected in fixed_shifts.items():
        actual = normalize(
            sheet.cell(
                row=row,
                column=column,
            ).value
        )

        if actual == expected:
            continue

        cell = sheet.cell(
            row=row,
            column=column,
        ).coordinate

        raise RuntimeError(
            "固定勤務が変更されました: "
            f"{cell} "
            f"({expected} -> {actual})"
        )


def numeric_to_half_units(
    value: int | float,
    label: str,
) -> int:
    """
    休日を0.5日単位の整数へ変換する。

    10日 -> 20
    0.5日 -> 1
    """
    scaled = float(value) * 2
    rounded = round(scaled)

    if abs(
        scaled - rounded
    ) > 1e-9:
        raise ValueError(
            f"{label} は0.5日単位で"
            f"入力してください: {value}"
        )

    return int(rounded)


def resolve_numeric_cell(
    formula_sheet: Worksheet,
    cached_sheet: Worksheet,
    coordinate: str,
    seen: set[str] | None = None,
) -> float:
    """数値または単純なセル参照式から値を取得する。"""
    seen = (
        set()
        if seen is None
        else seen
    )

    if coordinate in seen:
        raise ValueError(
            "循環参照を検出しました: "
            f"{coordinate}"
        )

    seen.add(coordinate)

    value = normalize(
        formula_sheet[
            coordinate
        ].value
    )

    if isinstance(
        value,
        (int, float),
    ):
        return float(value)

    if isinstance(value, str):
        match = re.fullmatch(
            r"=\$?([A-Z]{1,3})\$?(\d+)",
            value,
        )

        if match:
            reference = (
                f"{match.group(1)}"
                f"{match.group(2)}"
            )

            return resolve_numeric_cell(
                formula_sheet,
                cached_sheet,
                reference,
                seen,
            )

    cached = normalize(
        cached_sheet[
            coordinate
        ].value
    )

    if isinstance(
        cached,
        (int, float),
    ):
        return float(cached)

    raise ValueError(
        f"{coordinate} の数値を"
        "取得できません。"
        "Excelで一度再計算して保存してから"
        "実行してください。"
    )


def count_paid_leave_half_units(
    sheet: Worksheet,
    employee: Employee,
    columns: list[int],
) -> int:
    units = 0

    for column in columns:
        reason = normalize(
            sheet.cell(
                employee.row + 1,
                column,
            ).value
        )

        if reason == "年休":
            units += 2

        elif reason in {
            "年/",
            "/年",
        }:
            units += 1

    return units


def build_holiday_targets(
    formula_sheet: Worksheet,
    cached_sheet: Worksheet,
    employees: list[Employee],
    columns: list[int],
) -> dict[str, HolidayTarget]:
    public_column = find_header_column(
        formula_sheet,
        "公休数",
    )

    targets = {}

    for employee in employees:
        cell = formula_sheet.cell(
            employee.row,
            public_column,
        )

        public_days = resolve_numeric_cell(
            formula_sheet,
            cached_sheet,
            cell.coordinate,
        )

        targets[
            employee.name
        ] = HolidayTarget(
            public_half_units=(
                numeric_to_half_units(
                    public_days,
                    f"{cell.coordinate} 公休数",
                )
            ),
            paid_half_units=(
                count_paid_leave_half_units(
                    formula_sheet,
                    employee,
                    columns,
                )
            ),
        )

    return targets


def add_penalty(
    penalties: list[Penalty],
    label: str,
    variable: cp_model.IntVar,
    weight: int,
) -> None:
    penalties.append(
        Penalty(
            label,
            variable,
            weight,
        )
    )


def shortage(
    model: cp_model.CpModel,
    expression,
    target: int,
    upper_bound: int,
    name: str,
) -> cp_model.IntVar:
    variable = model.new_int_var(
        0,
        upper_bound,
        name,
    )

    model.add(
        variable
        >= target - expression
    )

    return variable


def excess(
    model: cp_model.CpModel,
    expression,
    limit: int,
    upper_bound: int,
    name: str,
) -> cp_model.IntVar:
    variable = model.new_int_var(
        0,
        upper_bound,
        name,
    )

    model.add(
        variable
        >= expression - limit
    )

    return variable


def build_model(
    sheet: Worksheet,
    employees: list[Employee],
    columns: list[int],
    holiday_targets: dict[
        str,
        HolidayTarget,
    ],
) -> ModelData:
    model = cp_model.CpModel()

    full_work = {}
    presence = {}
    off_half_units = {}

    fixed_shifts = {}
    existing_changes = []
    penalties = []

    for employee in employees:
        for column in columns:
            key = (
                employee.row,
                column,
            )

            work = normalize(
                sheet.cell(
                    employee.row,
                    column,
                ).value
            )

            reason = normalize(
                sheet.cell(
                    employee.row + 1,
                    column,
                ).value
            )

            work_var = model.new_bool_var(
                f"work_"
                f"{employee.row}_"
                f"{column}"
            )

            full_work[key] = work_var

            # 勤務と理由が両方入力済みなら絶対固定。
            if is_fixed_shift(
                work,
                reason,
            ):
                fixed_shifts[key] = work

                if work == "出":
                    model.add(
                        work_var == 1
                    )
                    presence[key] = 1
                    off_half_units[key] = 0

                elif work == "休":
                    model.add(
                        work_var == 0
                    )
                    presence[key] = 0
                    off_half_units[key] = 2

                else:
                    model.add(
                        work_var == 0
                    )
                    presence[key] = 1
                    off_half_units[key] = 1

                continue

            # 自動生成では半休を作らない。
            # 既にある半休はそのまま残す。
            if work in {
                "半/",
                "/半",
            }:
                fixed_shifts[key] = work

                model.add(
                    work_var == 0
                )

                presence[key] = 1
                off_half_units[key] = 1

                continue

            presence[key] = work_var

            # 出勤なら休日0、休みなら1日=2単位。
            off_half_units[key] = (
                2 - 2 * work_var
            )

            # 理由なしの既存予定も最優先で維持する。
            if work == "出":
                changed = model.new_bool_var(
                    f"change_existing_"
                    f"{employee.row}_"
                    f"{column}"
                )

                model.add(
                    changed + work_var == 1
                )

                existing_changes.append(
                    (
                        (
                            f"{employee.name} "
                            f"{day_label(sheet, column)}日 "
                            "既存「出」を変更"
                        ),
                        changed,
                    )
                )

            elif work == "休":
                changed = model.new_bool_var(
                    f"change_existing_"
                    f"{employee.row}_"
                    f"{column}"
                )

                model.add(
                    changed == work_var
                )

                existing_changes.append(
                    (
                        (
                            f"{employee.name} "
                            f"{day_label(sheet, column)}日 "
                            "既存「休」を変更"
                        ),
                        changed,
                    )
                )

    holiday_off = {}
    holiday_deviation = {}
    max_deviation_bound = 0

    for employee in employees:
        actual = sum(
            off_half_units[
                (
                    employee.row,
                    column,
                )
            ]
            for column in columns
        )

        target = (
            holiday_targets[
                employee.name
            ].total_half_units
        )

        upper_bound = (
            len(columns) * 2
            + target
        )

        deviation = model.new_int_var(
            0,
            upper_bound,
            (
                "holiday_deviation_"
                f"{employee.name}"
            ),
        )

        model.add_abs_equality(
            deviation,
            actual - target,
        )

        holiday_off[
            employee.name
        ] = actual

        holiday_deviation[
            employee.name
        ] = deviation

        max_deviation_bound = max(
            max_deviation_bound,
            upper_bound,
        )

    max_holiday_deviation = (
        model.new_int_var(
            0,
            max_deviation_bound,
            "max_holiday_deviation",
        )
    )

    model.add_max_equality(
        max_holiday_deviation,
        list(
            holiday_deviation.values()
        ),
    )

    # 終日の「出」が4日続くことを避ける。
    for employee in employees:
        for index in range(
            len(columns) - 3
        ):
            window = columns[
                index:index + 4
            ]

            if window != list(
                range(
                    window[0],
                    window[0] + 4,
                )
            ):
                continue

            work_count = sum(
                full_work[
                    (
                        employee.row,
                        column,
                    )
                ]
                for column in window
            )

            violation = excess(
                model,
                work_count,
                3,
                1,
                (
                    "four_consecutive_"
                    f"{employee.row}_"
                    f"{window[0]}"
                ),
            )

            add_penalty(
                penalties,
                (
                    f"{employee.name}: "
                    f"{day_label(sheet, window[0])}"
                    "〜"
                    f"{day_label(sheet, window[-1])}"
                    "日が4連続「出」"
                ),
                violation,
                WEIGHT_FOUR_CONSECUTIVE,
            )

    ots = [
        employee
        for employee in employees
        if employee.is_ot
    ]

    if not ots:
        raise ValueError(
            "作業療法士が"
            "A列に登録されていません。"
        )

    ward_members = {
        ward: [
            employee
            for employee in employees
            if employee.ward == ward
        ]
        for ward in WARD_MEMBERS
    }

    for column in columns:
        current_weekday = weekday(
            sheet,
            column,
        )

        day = day_label(
            sheet,
            column,
        )

        if current_weekday in MON_TO_SAT:
            for ward in WARD_MEMBERS:
                leaders = [
                    employee
                    for employee
                    in ward_members[ward]
                    if (
                        employee.is_dedicated
                        or employee.is_assigned
                    )
                ]

                leader_count = sum(
                    presence[
                        (
                            employee.row,
                            column,
                        )
                    ]
                    for employee in leaders
                )

                violation = shortage(
                    model,
                    leader_count,
                    1,
                    1,
                    (
                        "leader_shortage_"
                        f"{ward}_{column}"
                    ),
                )

                add_penalty(
                    penalties,
                    (
                        f"{day}日"
                        f"({current_weekday}) "
                        f"{ward}病棟 "
                        "専従・専任不在"
                    ),
                    violation,
                    WEIGHT_LEADER,
                )

            ot_count = sum(
                presence[
                    (
                        employee.row,
                        column,
                    )
                ]
                for employee in ots
            )

            ot_minimum = shortage(
                model,
                ot_count,
                1,
                1,
                f"ot_min_{column}",
            )

            ot_second = shortage(
                model,
                ot_count,
                2,
                2,
                f"ot_second_{column}",
            )

            add_penalty(
                penalties,
                (
                    f"{day}日"
                    f"({current_weekday}) "
                    "OTが0人"
                ),
                ot_minimum,
                WEIGHT_OT_MINIMUM,
            )

            add_penalty(
                penalties,
                (
                    f"{day}日"
                    f"({current_weekday}) "
                    "OTが2人未満"
                ),
                ot_second,
                WEIGHT_OT_SECOND,
            )

        if current_weekday == "日":
            for ward, members in (
                ward_members.items()
            ):
                count = sum(
                    presence[
                        (
                            employee.row,
                            column,
                        )
                    ]
                    for employee in members
                )

                difference = (
                    model.new_int_var(
                        0,
                        len(members),
                        (
                            f"sunday_"
                            f"{ward}_"
                            f"{column}"
                        ),
                    )
                )

                model.add_abs_equality(
                    difference,
                    count - 2,
                )

                add_penalty(
                    penalties,
                    (
                        f"{day}日(日) "
                        f"{ward}病棟 "
                        "2人配置との差"
                    ),
                    difference,
                    WEIGHT_SUNDAY_STAFF,
                )

        if current_weekday in {
            "水",
            "木",
            "金",
        }:
            for ward, members in (
                ward_members.items()
            ):
                working = sum(
                    presence[
                        (
                            employee.row,
                            column,
                        )
                    ]
                    for employee in members
                )

                off_count = (
                    len(members)
                    - working
                )

                over_two = excess(
                    model,
                    off_count,
                    2,
                    len(members),
                    (
                        f"off_over_two_"
                        f"{ward}_"
                        f"{column}"
                    ),
                )

                over_three = excess(
                    model,
                    off_count,
                    3,
                    len(members),
                    (
                        f"off_over_three_"
                        f"{ward}_"
                        f"{column}"
                    ),
                )

                add_penalty(
                    penalties,
                    (
                        f"{day}日"
                        f"({current_weekday}) "
                        f"{ward}病棟 "
                        "休み2人超過"
                    ),
                    over_two,
                    WEIGHT_THIRD_OFF,
                )

                add_penalty(
                    penalties,
                    (
                        f"{day}日"
                        f"({current_weekday}) "
                        f"{ward}病棟 "
                        "休み3人超過"
                    ),
                    over_three,
                    WEIGHT_TOO_MANY_OFF,
                )

        if current_weekday in {
            "月",
            "火",
            "土",
        }:
            counts = {
                ward: sum(
                    presence[
                        (
                            employee.row,
                            column,
                        )
                    ]
                    for employee in members
                )
                for ward, members
                in ward_members.items()
            }

            ward1_few = shortage(
                model,
                counts[1],
                2,
                len(
                    ward_members[1]
                ),
                f"ward1_few_{column}",
            )

            ward1_many = excess(
                model,
                counts[1],
                3,
                len(
                    ward_members[1]
                ),
                f"ward1_many_{column}",
            )

            add_penalty(
                penalties,
                (
                    f"{day}日"
                    f"({current_weekday}) "
                    "1病棟 2人未満"
                ),
                ward1_few,
                WEIGHT_WARD1,
            )

            add_penalty(
                penalties,
                (
                    f"{day}日"
                    f"({current_weekday}) "
                    "1病棟 3人超過"
                ),
                ward1_many,
                WEIGHT_WARD1,
            )

            ward2_difference = (
                model.new_int_var(
                    0,
                    len(
                        ward_members[2]
                    ),
                    (
                        "ward2_diff_"
                        f"{column}"
                    ),
                )
            )

            model.add_abs_equality(
                ward2_difference,
                counts[2] - 4,
            )

            add_penalty(
                penalties,
                (
                    f"{day}日"
                    f"({current_weekday}) "
                    "2病棟 4人配置との差"
                ),
                ward2_difference,
                WEIGHT_WARD2,
            )

            ward3_few = shortage(
                model,
                counts[3],
                3,
                len(
                    ward_members[3]
                ),
                f"ward3_few_{column}",
            )

            add_penalty(
                penalties,
                (
                    f"{day}日"
                    f"({current_weekday}) "
                    "3病棟 3人未満"
                ),
                ward3_few,
                WEIGHT_WARD3,
            )

    # 「休→出→休」は禁止しない。
    # 同条件なら少しだけ避ける。
    for employee in employees:
        for index in range(
            1,
            len(columns) - 1,
        ):
            previous_column = (
                columns[index - 1]
            )
            column = columns[index]
            next_column = (
                columns[index + 1]
            )

            if (
                next_column
                - previous_column
                != 2
            ):
                continue

            isolated = (
                model.new_bool_var(
                    (
                        f"isolated_"
                        f"{employee.row}_"
                        f"{column}"
                    )
                )
            )

            model.add(
                isolated
                >= (
                    full_work[
                        (
                            employee.row,
                            column,
                        )
                    ]
                    - full_work[
                        (
                            employee.row,
                            previous_column,
                        )
                    ]
                    - full_work[
                        (
                            employee.row,
                            next_column,
                        )
                    ]
                )
            )

            add_penalty(
                penalties,
                (
                    f"{employee.name}: "
                    f"{day_label(sheet, column)}日"
                    "が1日だけ出勤"
                ),
                isolated,
                WEIGHT_ISOLATED_WORK,
            )

    return ModelData(
        model=model,
        full_work=full_work,
        fixed_shifts=fixed_shifts,
        existing_changes=(
            existing_changes
        ),
        penalties=penalties,
        holiday_targets=(
            holiday_targets
        ),
        holiday_off=holiday_off,
        holiday_deviation=(
            holiday_deviation
        ),
        max_holiday_deviation=(
            max_holiday_deviation
        ),
    )


def new_solver(
    max_seconds: float = 60,
) -> cp_model.CpSolver:
    solver = cp_model.CpSolver()

    solver.parameters.max_time_in_seconds = (
        max_seconds
    )

    # 同じ入力から同じ結果を得やすくする。
    solver.parameters.random_seed = 42
    solver.parameters.num_search_workers = 1

    return solver


def solve_stage(
    model: cp_model.CpModel,
    objective,
    stage_name: str,
) -> tuple[
    cp_model.CpSolver,
    int,
]:
    model.minimize(objective)

    solver = new_solver()
    status = solver.solve(model)

    if status not in {
        cp_model.OPTIMAL,
        cp_model.FEASIBLE,
    }:
        raise RuntimeError(
            f"{stage_name}で"
            "勤務表を作成できませんでした。"
        )

    return (
        solver,
        int(
            solver.value(
                objective
            )
        ),
    )


def solve_schedule(
    data: ModelData,
) -> cp_model.CpSolver:
    """
    優先順位を段階的に固定しながら最適化する。

    1. 既存予定
    2. 個人休日数の最大ズレ
    3. 全員の休日数ズレ
    4. その他の勤務条件
    """
    model = data.model

    if data.existing_changes:
        existing_total = sum(
            variable
            for _, variable
            in data.existing_changes
        )

        _, best = solve_stage(
            model,
            existing_total,
            "既存予定の維持",
        )

        model.add(
            existing_total == best
        )

    _, best = solve_stage(
        model,
        data.max_holiday_deviation,
        "休日数の公平化",
    )

    model.add(
        data.max_holiday_deviation
        == best
    )

    holiday_total = sum(
        data.holiday_deviation.values()
    )

    _, best = solve_stage(
        model,
        holiday_total,
        "休日数の調整",
    )

    model.add(
        holiday_total == best
    )

    other_cost = sum(
        penalty.variable
        * penalty.weight
        for penalty in data.penalties
    )

    solver, _ = solve_stage(
        model,
        other_cost,
        "勤務条件の最適化",
    )

    return solver


def write_schedule(
    sheet: Worksheet,
    employees: list[Employee],
    columns: list[int],
    solver: cp_model.CpSolver,
    data: ModelData,
) -> None:
    """
    自動配置した通常出勤は空白にする。

    休み       -> 「休」
    通常出勤   -> 空白
    既存の「出」-> 維持
    固定勤務   -> 変更しない
    """
    for employee in employees:
        for column in columns:
            key = (
                employee.row,
                column,
            )

            if key in data.fixed_shifts:
                continue

            cell = sheet.cell(
                employee.row,
                column,
            )

            original = normalize(
                cell.value
            )

            is_working = (
                solver.value(
                    data.full_work[key]
                )
                == 1
            )

            if is_working:
                cell.value = (
                    "出"
                    if original == "出"
                    else None
                )
            else:
                cell.value = "休"


def request_recalculation(
    workbook,
) -> None:
    calculation = getattr(
        workbook,
        "calculation",
        None,
    )

    if calculation is None:
        return

    calculation.calcMode = "auto"
    calculation.fullCalcOnLoad = True
    calculation.forceFullCalc = True
    calculation.calcOnSave = True


def save_safely(
    workbook,
    formulas: dict[str, str],
    fixed_shifts: dict[
        tuple[int, int],
        str,
    ],
) -> None:
    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if TEMP_FILE.exists():
        TEMP_FILE.unlink()

    request_recalculation(
        workbook
    )

    workbook.save(
        TEMP_FILE
    )

    check_book = load_workbook(
        TEMP_FILE,
        data_only=False,
    )

    try:
        check_sheet = check_book[
            SHEET_NAME
        ]

        validate_formulas(
            check_sheet,
            formulas,
        )

        validate_fixed_shifts(
            check_sheet,
            fixed_shifts,
        )

    finally:
        check_book.close()

    try:
        TEMP_FILE.replace(
            OUTPUT_FILE
        )

    except PermissionError as exc:
        raise PermissionError(
            "kinmu_output.xlsx を"
            "Excelで開いている可能性があります。"
            "閉じてからもう一度"
            "実行してください。"
        ) from exc


def half_units_to_days(
    value: int,
) -> float:
    return value / 2


def print_result(
    solver: cp_model.CpSolver,
    employees: list[Employee],
    data: ModelData,
) -> None:
    print(
        "\n===== 個人別 休日数 ====="
    )

    for employee in employees:
        name = employee.name

        target = (
            data.holiday_targets[
                name
            ]
        )

        actual = int(
            solver.value(
                data.holiday_off[
                    name
                ]
            )
        )

        difference = (
            target.total_half_units
            - actual
        )

        if difference == 0:
            result = "OK"

        elif difference > 0:
            result = (
                "あと"
                f"{half_units_to_days(difference):g}"
                "日必要"
            )

        else:
            result = (
                f"{half_units_to_days(-difference):g}"
                "日多い"
            )

        print(
            f"{name}: "
            f"必要"
            f"{half_units_to_days(target.total_half_units):g}"
            "日 / "
            f"割当"
            f"{half_units_to_days(actual):g}"
            "日 -> "
            f"{result}"
        )

    changed = [
        label
        for label, variable
        in data.existing_changes
        if solver.value(variable) > 0
    ]

    if changed:
        print(
            "\n===== 既存予定の変更 ====="
        )

        for label in changed:
            print(
                f"⚠ {label}"
            )

    violations = [
        (
            penalty.weight,
            int(
                solver.value(
                    penalty.variable
                )
            ),
            penalty.label,
        )
        for penalty in data.penalties
        if solver.value(
            penalty.variable
        ) > 0
    ]

    if violations:
        print(
            "\n===== "
            "満たせなかった希望条件 "
            "====="
        )

        for (
            _,
            value,
            label,
        ) in sorted(
            violations,
            reverse=True,
        ):
            print(
                f"⚠ {label}"
                f"（差: {value}）"
            )


def main() -> None:
    if not INPUT_FILE.exists():
        raise FileNotFoundError(
            "勤務表が見つかりません: "
            f"{INPUT_FILE}"
        )

    workbook = load_workbook(
        INPUT_FILE,
        data_only=False,
    )

    cached_workbook = load_workbook(
        INPUT_FILE,
        data_only=True,
    )

    try:
        if (
            SHEET_NAME
            not in workbook.sheetnames
        ):
            raise ValueError(
                f"シート「{SHEET_NAME}」"
                "がありません。"
            )

        sheet = workbook[
            SHEET_NAME
        ]

        cached_sheet = (
            cached_workbook[
                SHEET_NAME
            ]
        )

        employees = discover_employees(
            sheet
        )

        columns = day_columns(
            sheet
        )

        errors = validate_input(
            sheet,
            employees,
            columns,
        )

        if errors:
            print(
                "\n===== 入力エラー ====="
            )

            for error in errors:
                print(
                    f"❌ {error}"
                )

            return

        formulas = capture_formulas(
            sheet
        )

        holiday_targets = (
            build_holiday_targets(
                sheet,
                cached_sheet,
                employees,
                columns,
            )
        )

        data = build_model(
            sheet,
            employees,
            columns,
            holiday_targets,
        )

        solver = solve_schedule(
            data
        )

        write_schedule(
            sheet,
            employees,
            columns,
            solver,
            data,
        )

        validate_formulas(
            sheet,
            formulas,
        )

        validate_fixed_shifts(
            sheet,
            data.fixed_shifts,
        )

        save_safely(
            workbook,
            formulas,
            data.fixed_shifts,
        )

        print(
            "\n勤務表を作成しました:"
        )

        print(
            OUTPUT_FILE
        )

        print_result(
            solver,
            employees,
            data,
        )

    finally:
        cached_workbook.close()
        workbook.close()


if __name__ == "__main__":
    main()