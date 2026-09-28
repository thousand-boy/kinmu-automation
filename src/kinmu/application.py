"""読み込み、検証、最適化、保存をつなぐ実行処理。"""

from __future__ import annotations

import argparse
from contextlib import closing
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import load_workbook

from .constants import MASTER_SHEET_NAME, SCHEDULE_SHEET_NAME
from .models import CellKey, Employee, HolidayTarget, SolveResult
from .optimizer import solve_schedule
from .reporting import print_result
from .roster import (
    discover_employees,
    load_employee_master,
    validate_employee_structure,
)
from .validation import (
    build_holiday_targets,
    capture_formulas,
    validate_fixed_shifts,
    validate_formulas,
    validate_holiday_feasibility,
    validate_input,
    validate_sunday_fixed_staffing,
    validate_written_holidays,
)
from .workbook import (
    day_columns,
    request_recalculation,
    synchronize_employee_rows,
    update_holiday_summary_formulas,
    write_schedule,
)

BASE_DIR = Path(__file__).resolve().parents[2]


def save_safely(
    workbook,
    formulas: dict[tuple[str, str], str],
    fixed_shifts: dict[CellKey, str],
    employees: list[Employee],
    columns: list[int],
    holiday_targets: dict[str, HolidayTarget],
    output_file: Path,
) -> None:
    """同じ出力先の一時ファイルを再検証してから、正式ファイルへ置換する。"""
    output_file = Path(output_file)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(
        dir=output_file.parent,
        prefix=f".{output_file.stem}-",
        suffix=".xlsx",
        delete=False,
    ) as temporary:
        temporary_path = Path(temporary.name)

    try:
        request_recalculation(workbook)
        workbook.save(temporary_path)
        with closing(load_workbook(temporary_path, data_only=False)) as saved:
            validate_formulas(saved, formulas)
            sheet = saved[SCHEDULE_SHEET_NAME]
            validate_fixed_shifts(sheet, fixed_shifts)
            validate_written_holidays(
                sheet, employees, columns, holiday_targets,
            )
            load_employee_master(saved[MASTER_SHEET_NAME])

        # 検証が完了するまで前回の出力を残す。置換は同じファイルシステム内。
        try:
            temporary_path.replace(output_file)
        except PermissionError as exc:
            raise PermissionError(
                f"{output_file.name} を置き換えられません。"
                "Excelで開いている場合は閉じ、保存先の権限も確認してください。"
            ) from exc
    finally:
        temporary_path.unlink(missing_ok=True)


def run(input_file: Path, output_file: Path) -> SolveResult:
    """入力ファイルを変更せず、勤務表を別ファイルへ出力する。"""
    input_file, output_file = Path(input_file), Path(output_file)
    if input_file.resolve() == output_file.resolve():
        raise ValueError("入力ファイルと出力ファイルには別のパスを指定してください。")
    if not input_file.is_file():
        raise FileNotFoundError(f"勤務表が見つかりません: {input_file}")

    with (
        closing(load_workbook(input_file, data_only=False)) as workbook,
        closing(load_workbook(input_file, data_only=True)) as cached_workbook,
    ):
        required = {SCHEDULE_SHEET_NAME, MASTER_SHEET_NAME}
        missing = required - set(workbook.sheetnames)
        if missing:
            raise ValueError(f"必要なシートがありません: {', '.join(sorted(missing))}")

        sheet = workbook[SCHEDULE_SHEET_NAME]
        cached_sheet = cached_workbook[SCHEDULE_SHEET_NAME]
        records = load_employee_master(workbook[MASTER_SHEET_NAME])
        added = synchronize_employee_rows(sheet, records)
        # 行挿入後も、既存職員の公休数キャッシュを同じ行から取得する。
        synchronize_employee_rows(cached_sheet, records)
        employees = discover_employees(sheet, records)
        validate_employee_structure(employees)
        columns = day_columns(sheet)

        errors = validate_input(sheet, employees, columns)
        if errors:
            raise ValueError("入力エラー:\n" + "\n".join(errors))
        targets = build_holiday_targets(
            sheet, cached_sheet, employees, columns,
        )
        validate_holiday_feasibility(sheet, employees, columns, targets)
        validate_sunday_fixed_staffing(sheet, employees, columns)
        update_holiday_summary_formulas(sheet, employees, columns)
        formulas = capture_formulas(workbook)

        result = solve_schedule(sheet, employees, columns, targets)
        write_schedule(sheet, employees, columns, result.solver, result.data)
        validate_formulas(workbook, formulas)
        validate_fixed_shifts(sheet, result.data.fixed_shifts)
        validate_written_holidays(sheet, employees, columns, targets)
        save_safely(
            workbook, formulas, result.data.fixed_shifts,
            employees, columns, targets, output_file,
        )

    print(f"\n職員マスタ: 在籍 {len(employees)}名")
    if added:
        print(f"Sheet1へ追加: {', '.join(added)}")
    print(f"勤務表を作成しました:\n{output_file}")
    print_result(result, employees)
    return result


def main(argv=None) -> None:
    parser = argparse.ArgumentParser(description="Excel勤務表を自動調整します。")
    parser.add_argument(
        "--input", type=Path, default=BASE_DIR / "input" / "kinmu_sample.xlsx",
        help="入力Excel（既定: input/kinmu_sample.xlsx）",
    )
    parser.add_argument(
        "--output", type=Path,
        default=BASE_DIR / "output" / "kinmu_output.xlsx",
        help="出力Excel（既定: output/kinmu_output.xlsx）",
    )
    args = parser.parse_args(argv)
    run(args.input, args.output)
