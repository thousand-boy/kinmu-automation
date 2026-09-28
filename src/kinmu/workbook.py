"""勤務表の行同期、数式・書式の引き継ぎ、計算結果の書き込み。"""

from __future__ import annotations

import re
from copy import copy, deepcopy

from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
from openpyxl.worksheet.worksheet import Worksheet
from ortools.sat.python import cp_model

from .constants import (
    DATE_ROW,
    DAY_END_COLUMN,
    DAY_START_COLUMN,
    WARD_STAFF_LIMITS,
    WEEKDAY_ROW,
    WEEKDAYS,
)
from .models import Employee, EmployeeMaster, ModelData
from .roster import _staff_rows, employee_display_label, normalize


_CELL_REFERENCE = re.compile(
    r"^(?:(?P<sheet>'(?:[^']|'')+'|[^!]+)!)?"
    r"(?P<start>\$?[A-Z]{1,3}\$?\d+)"
    r"(?::(?P<end>\$?[A-Z]{1,3}\$?\d+))?$"
)


def _sheet_reference_name(value: str) -> str:
    if value.startswith("'") and value.endswith("'"):
        return value[1:-1].replace("''", "'")
    return value


def _split_cell_reference(value: str) -> tuple[str, int]:
    match = re.fullmatch(r"(\$?[A-Z]{1,3})(\$?\d+)", value)
    if match is None:
        raise ValueError(f"セル参照を解析できません: {value}")
    return match.group(1), int(match.group(2).replace("$", ""))


def _join_cell_reference(column: str, original: str, row: int) -> str:
    row_prefix = "$" if "$" in original[len(column):] else ""
    return f"{column}{row_prefix}{row}"


def _shift_formula_range(
    value: str,
    *,
    formula_sheet_name: str,
    edited_sheet_name: str,
    insert_at: int,
    amount: int,
    expand_at_boundary: bool = False,
) -> str:
    match = _CELL_REFERENCE.fullmatch(value)
    if match is None:
        return value

    sheet_name = match.group("sheet")
    referenced_sheet = (
        _sheet_reference_name(sheet_name) if sheet_name else formula_sheet_name
    )
    if referenced_sheet != edited_sheet_name:
        return value

    start, end = match.group("start", "end")
    start_column, start_row = _split_cell_reference(start)
    if end is None:
        if start_row >= insert_at:
            start_row += amount
        updated = _join_cell_reference(start_column, start, start_row)
    else:
        end_column, end_row = _split_cell_reference(end)
        if insert_at <= start_row:
            start_row += amount
            end_row += amount
        elif insert_at <= end_row or (
            expand_at_boundary and insert_at == end_row + 1
        ):
            # 境界直後の行を取り込むのは病棟集計だけ。個人の範囲は広げない。
            end_row += amount
        updated = (
            f"{_join_cell_reference(start_column, start, start_row)}:"
            f"{_join_cell_reference(end_column, end, end_row)}"
        )
    return f"{sheet_name}!{updated}" if sheet_name else updated


def _shift_formula_rows(
    formula: str,
    *,
    formula_sheet_name: str,
    edited_sheet_name: str,
    insert_at: int,
    amount: int,
    expand_at_boundary: bool = False,
) -> str:
    """文字列リテラルを変えず、A1形式の直接参照だけを行挿入に追従させる。"""
    tokenizer = Tokenizer(formula)
    for token in tokenizer.items:
        if token.type == "OPERAND" and token.subtype == "RANGE":
            token.value = _shift_formula_range(
                token.value,
                formula_sheet_name=formula_sheet_name,
                edited_sheet_name=edited_sheet_name,
                insert_at=insert_at,
                amount=amount,
                expand_at_boundary=expand_at_boundary,
            )
    return tokenizer.render()


def _shift_formula_expression(value, **kwargs):
    """入力規則・条件付き書式では数式の先頭の = が省略される。"""
    if not isinstance(value, str) or not value:
        return value
    has_equals = value.startswith("=")
    formula = value if has_equals else f"={value}"
    shifted = _shift_formula_rows(formula, **kwargs)
    return shifted if has_equals else shifted[1:]


def _shift_cell_range(
    cell_range: CellRange, insert_at: int, amount: int
) -> CellRange:
    updated = copy(cell_range)
    if updated.min_row >= insert_at:
        updated.shift(row_shift=amount)
    elif updated.max_row >= insert_at:
        updated.max_row += amount
    return updated


def _shift_row_dimensions(
    sheet: Worksheet, insert_at: int, amount: int
) -> None:
    dimensions = {
        row: copy(dimension)
        for row, dimension in sheet.row_dimensions.items()
        if row >= insert_at
    }
    for row in dimensions:
        del sheet.row_dimensions[row]
    for row, dimension in dimensions.items():
        dimension.index = row + amount
        sheet.row_dimensions[row + amount] = dimension


def _shift_sheet_features(
    sheet: Worksheet, insert_at: int, amount: int
) -> None:
    """insert_rows が更新しない結合、行属性、規則、直接参照を補完する。"""
    merged_ranges = list(sheet.merged_cells.ranges)
    for cell_range in merged_ranges:
        sheet.unmerge_cells(str(cell_range))

    validations = list(sheet.data_validations.dataValidation)
    conditional_rules = [
        (list(conditional.sqref.ranges), deepcopy(conditional.rules))
        for conditional in sheet.conditional_formatting
    ]
    for conditional in list(sheet.conditional_formatting):
        del sheet.conditional_formatting[str(conditional.sqref)]

    sheet.insert_rows(insert_at, amount)
    _shift_row_dimensions(sheet, insert_at, amount)

    for cell_range in merged_ranges:
        shifted = _shift_cell_range(cell_range, insert_at, amount)
        sheet.merge_cells(str(shifted))

    formula_options = {
        "formula_sheet_name": sheet.title,
        "edited_sheet_name": sheet.title,
        "insert_at": insert_at,
        "amount": amount,
    }
    for validation in validations:
        validation.sqref = MultiCellRange(
            ranges={
                _shift_cell_range(item, insert_at, amount)
                for item in validation.ranges.ranges
            }
        )
        validation.formula1 = _shift_formula_expression(
            validation.formula1, **formula_options
        )
        validation.formula2 = _shift_formula_expression(
            validation.formula2, **formula_options
        )

    for ranges, rules in conditional_rules:
        shifted_ranges = " ".join(
            str(_shift_cell_range(item, insert_at, amount)) for item in ranges
        )
        for rule in rules:
            if rule.formula:
                rule.formula = [
                    _shift_formula_expression(formula, **formula_options)
                    for formula in rule.formula
                ]
            sheet.conditional_formatting.add(shifted_ranges, rule)

    for formula_sheet in sheet.parent.worksheets:
        for row in formula_sheet.iter_rows():
            for cell in row:
                if cell.data_type != "f":
                    continue
                cell.value = _shift_formula_rows(
                    cell.value,
                    formula_sheet_name=formula_sheet.title,
                    edited_sheet_name=sheet.title,
                    insert_at=insert_at,
                    amount=amount,
                    expand_at_boundary=(
                        formula_sheet is sheet
                        and normalize(sheet.cell(cell.row, 2).value) == "階休み"
                    ),
                )


def _remap_template_formula(
    formula: str, *, source_row: int, target_row: int, sheet_name: str
) -> str:
    """雛形の2行と一致する行参照を移し、それ以外の参照は保つ。"""
    tokenizer = Tokenizer(formula)
    row_map = {source_row: target_row, source_row + 1: target_row + 1}
    for token in tokenizer.items:
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        match = _CELL_REFERENCE.fullmatch(token.value)
        if match is None:
            continue
        referenced_sheet = match.group("sheet")
        if (
            referenced_sheet
            and _sheet_reference_name(referenced_sheet) != sheet_name
        ):
            continue

        pieces = []
        for reference in match.group("start", "end"):
            if reference is None:
                continue
            column, row = _split_cell_reference(reference)
            pieces.append(
                _join_cell_reference(column, reference, row_map.get(row, row))
            )
        replacement = ":".join(pieces)
        if referenced_sheet:
            replacement = f"{referenced_sheet}!{replacement}"
        token.value = replacement
    return tokenizer.render()


def _remap_template_formula_expression(value, **kwargs):
    if not isinstance(value, str) or not value:
        return value
    has_equals = value.startswith("=")
    formula = value if has_equals else f"={value}"
    remapped = _remap_template_formula(formula, **kwargs)
    return remapped if has_equals else remapped[1:]


def _copy_row_pair(sheet: Worksheet, source_row: int, target_row: int) -> None:
    for offset in (0, 1):
        dimension = copy(sheet.row_dimensions[source_row + offset])
        dimension.index = target_row + offset
        sheet.row_dimensions[target_row + offset] = dimension
        for column in range(1, sheet.max_column + 1):
            source = sheet.cell(source_row + offset, column)
            target = sheet.cell(target_row + offset, column)
            # StyleArray の複製で罫線・配置・表示形式・保護をまとめて引き継ぐ。
            target._style = copy(source._style)
            target.comment = copy(source.comment)
            target.hyperlink = copy(source.hyperlink)
            value = source.value
            if source.data_type == "f":
                value = _remap_template_formula(
                    value,
                    source_row=source_row,
                    target_row=target_row,
                    sheet_name=sheet.title,
                )
            target.value = value

    sheet.merge_cells(
        start_row=target_row,
        start_column=1,
        end_row=target_row + 1,
        end_column=1,
    )


def _copied_rule_range(
    cell_range: CellRange, source_row: int, target_row: int
) -> str | None:
    first = max(cell_range.min_row, source_row)
    last = min(cell_range.max_row, source_row + 1)
    if first > last:
        return None
    return str(
        CellRange(
            min_col=cell_range.min_col,
            max_col=cell_range.max_col,
            min_row=target_row + first - source_row,
            max_row=target_row + last - source_row,
        )
    )


def _copy_row_rules(
    sheet: Worksheet, source_row: int, target_row: int
) -> None:
    for validation in sheet.data_validations.dataValidation:
        for item in list(validation.ranges.ranges):
            destination = _copied_rule_range(item, source_row, target_row)
            if destination:
                validation.add(destination)

    additions = []
    for conditional in sheet.conditional_formatting:
        for item in conditional.sqref.ranges:
            destination = _copied_rule_range(item, source_row, target_row)
            if destination:
                additions.append((destination, deepcopy(conditional.rules)))

    for destination, rules in additions:
        for rule in rules:
            if rule.formula:
                rule.formula = [
                    _remap_template_formula_expression(
                        formula,
                        source_row=source_row,
                        target_row=target_row,
                        sheet_name=sheet.title,
                    )
                    for formula in rule.formula
                ]
            sheet.conditional_formatting.add(destination, rule)


def _ward_template_and_insertion_row(
    sheet: Worksheet, master_records: list[EmployeeMaster], ward: int
) -> tuple[int, int]:
    rows = _staff_rows(sheet, master_records)
    ward_ids = {
        record.employee_id for record in master_records if record.ward == ward
    }
    ward_rows = [
        row for employee_id, row in rows.items() if employee_id in ward_ids
    ]
    if not ward_rows:
        raise ValueError(f"{ward}病棟に、追加時の雛形として使える既存職員がいません。")

    # 雛形を固定し、複数人追加時に条件付き書式の複製が連鎖しないようにする。
    template_row = min(ward_rows)
    insertion_row = max(ward_rows) + 2
    while insertion_row <= sheet.max_row:
        row_type = normalize(sheet.cell(insertion_row, 2).value)
        if row_type == "階休み":
            break
        if row_type == "勤務希望":
            raise ValueError(
                f"{ward}病棟の職員範囲を特定できません。"
                f" Sheet1!B{insertion_row}を確認してください。"
            )
        insertion_row += 1
    return template_row, insertion_row


def synchronize_employee_rows(
    sheet: Worksheet, master_records: list[EmployeeMaster]
) -> list[str]:
    """在籍職員の不足行を病棟内へ追加し、退職者の行と入力は残す。"""
    added = []
    for record in master_records:
        rows = _staff_rows(sheet, master_records)
        if record.employee_id in rows:
            if record.is_active:
                label_cell = sheet.cell(rows[record.employee_id], 1)
                label_cell.value = employee_display_label(record)
            continue
        if not record.is_active:
            continue

        template_row, insertion_row = _ward_template_and_insertion_row(
            sheet, master_records, record.ward
        )
        _shift_sheet_features(sheet, insertion_row, 2)
        _copy_row_pair(sheet, template_row, insertion_row)
        _copy_row_rules(sheet, template_row, insertion_row)
        for row in (insertion_row, insertion_row + 1):
            for column in range(DAY_START_COLUMN, DAY_END_COLUMN + 1):
                sheet.cell(row, column).value = None

        sheet.cell(insertion_row, 1).value = employee_display_label(record)
        sheet.cell(insertion_row, 2).value = "勤務希望"
        sheet.cell(insertion_row + 1, 2).value = "理由欄"
        added.append(record.display_name)
    return added


def update_holiday_summary_formulas(
    sheet: Worksheet, employees: list[Employee], columns: list[int]
) -> None:
    """病棟順に並ぶ「階休み」と「全体休み」を在籍職員だけの集計にする。"""
    wards = sorted(WARD_STAFF_LIMITS)
    ward_rows = []
    overall_rows = []
    for row in range(1, sheet.max_row + 1):
        label = normalize(sheet.cell(row, 2).value)
        if label == "階休み":
            ward_rows.append(row)
        elif label == "全体休み":
            overall_rows.append(row)

    if len(ward_rows) != len(wards) or len(overall_rows) != 1:
        raise ValueError(
            "勤務表の休み集計行を確認してください。"
            f" 病棟順の「階休み」が{len(wards)}行、"
            "「全体休み」が1行必要です。"
        )
    if {employee.ward for employee in employees} - set(wards):
        raise ValueError("休み集計に未対応の病棟があります。")

    employees_by_ward = {
        ward: sorted(
            employee.row for employee in employees if employee.ward == ward
        )
        for ward in wards
    }
    for column in columns:
        letter = get_column_letter(column)
        for ward, summary_row in zip(wards, ward_rows, strict=True):
            terms = [
                f'COUNTIF({letter}{row},"休")'
                f'+COUNTIF({letter}{row},"半/")/2'
                f'+COUNTIF({letter}{row},"/半")/2'
                for row in employees_by_ward[ward]
            ]
            sheet.cell(summary_row, column).value = (
                f"=SUM({','.join(terms)})" if terms else "=0"
            )
        references = ",".join(f"{letter}{row}" for row in ward_rows)
        sheet.cell(overall_rows[0], column).value = f"=SUM({references})"


def day_columns(sheet: Worksheet) -> list[int]:
    columns = []
    for column in range(DAY_START_COLUMN, DAY_END_COLUMN + 1):
        date_value = normalize(sheet.cell(DATE_ROW, column).value)
        weekday_value = normalize(sheet.cell(WEEKDAY_ROW, column).value)
        if date_value is None and weekday_value is None:
            continue
        if weekday_value not in WEEKDAYS:
            coordinate = sheet.cell(WEEKDAY_ROW, column).coordinate
            raise ValueError(f"{coordinate} の曜日を認識できません: {weekday_value}")
        columns.append(column)
    if not columns:
        raise ValueError("勤務対象の日付列が見つかりません。")
    return columns


def day_label(sheet: Worksheet, column: int) -> str:
    value = normalize(sheet.cell(DATE_ROW, column).value)
    return str(value if value is not None else column - DAY_START_COLUMN + 1)


def weekday(sheet: Worksheet, column: int) -> str:
    value = normalize(sheet.cell(WEEKDAY_ROW, column).value)
    if value not in WEEKDAYS:
        raise ValueError(f"曜日を認識できません: {value}")
    return value


def find_header_column(sheet: Worksheet, header: str) -> int:
    for row in sheet.iter_rows(min_row=1, max_row=WEEKDAY_ROW):
        for cell in row:
            if normalize(cell.value) == header:
                return cell.column
    raise ValueError(f"見出し「{header}」が見つかりません。")


def write_schedule(
    sheet: Worksheet,
    employees: list[Employee],
    columns: list[int],
    solver: cp_model.CpSolver,
    data: ModelData,
) -> None:
    """固定勤務を保護し、自動勤務は空欄（元の「出」は維持）、休みは「休」にする。"""
    for employee in employees:
        for column in columns:
            key = employee.row, column
            if key in data.fixed_shifts:
                continue
            cell = sheet.cell(employee.row, column)
            if solver.value(data.full_work[key]) == 1:
                cell.value = "出" if normalize(cell.value) == "出" else None
            else:
                cell.value = "休"


def request_recalculation(workbook) -> None:
    calculation = workbook.calculation
    if calculation is not None:
        calculation.calcMode = "auto"
        calculation.fullCalcOnLoad = True
        calculation.forceFullCalc = True
        calculation.calcOnSave = True
