"""入力値・休日数・出力の不変条件の検証。"""

from __future__ import annotations

import math
import re

from openpyxl.worksheet.worksheet import Worksheet

from .constants import (
    FIXED_REASON_VALUES,
    FIXED_WORK_VALUES,
    REASON_VALUES,
    WORK_VALUES,
)
from .models import CellKey, Employee, HolidayTarget
from .roster import normalize
from .workbook import day_label, find_header_column, weekday


def is_fixed_shift(work, reason) -> bool:
    return work in FIXED_WORK_VALUES and reason in FIXED_REASON_VALUES


def validate_input(
    sheet: Worksheet, employees: list[Employee], columns: list[int],
) -> list[str]:
    errors = []
    for employee in employees:
        for column in columns:
            work_cell = sheet.cell(employee.row, column)
            reason_cell = sheet.cell(employee.row + 1, column)
            work = normalize(work_cell.value)
            reason = normalize(reason_cell.value)
            address = work_cell.coordinate
            if work not in WORK_VALUES:
                errors.append(f"{address}: 不明な勤務「{work}」")
            if reason not in REASON_VALUES:
                errors.append(f"{reason_cell.coordinate}: 不明な理由「{reason}」")
            if reason in FIXED_REASON_VALUES and work not in FIXED_WORK_VALUES:
                errors.append(
                    f"{address}: 理由「{reason}」がありますが勤務が未入力です。"
                )
            if reason == "年休" and work != "休":
                errors.append(f"{address}: 年休は「休」と組み合わせてください。")
            if reason in {"年/", "/年"} and work not in {"半/", "/半"}:
                errors.append(
                    f"{address}: 半日の年休は「半/」または「/半」と"
                    "組み合わせてください。"
                )
    return errors


def numeric_to_half_units(value: int | float, label: str) -> int:
    """CP-SATの整数制約で扱えるよう、1日を2単位に換算する。"""
    if isinstance(value, bool) or not math.isfinite(value) or value < 0:
        raise ValueError(f"{label} は0以上の有限な日数を入力してください: {value}")
    scaled = value * 2
    if not math.isfinite(scaled):
        raise ValueError(f"{label} の日数が大きすぎます: {value}")
    rounded = round(scaled)
    if abs(scaled - rounded) > 1e-9:
        raise ValueError(f"{label} は0.5日単位で入力してください: {value}")
    return rounded


def resolve_numeric_cell(
    formula_sheet: Worksheet,
    cached_sheet: Worksheet,
    coordinate: str,
    seen: set[str] | None = None,
) -> float:
    """数値・単純セル参照・Excelの計算済みキャッシュの順に解決する。"""
    visited = set() if seen is None else set(seen)
    while coordinate not in visited:
        visited.add(coordinate)
        value = normalize(formula_sheet[coordinate].value)
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return float(value)
        if isinstance(value, str):
            match = re.fullmatch(r"=\$?([A-Z]{1,3})\$?(\d+)", value)
            if match:
                coordinate = "".join(match.groups())
                continue
        cached = normalize(cached_sheet[coordinate].value)
        if isinstance(cached, (int, float)) and not isinstance(cached, bool):
            return float(cached)
        raise ValueError(
            f"{coordinate} の数値を取得できません。"
            "Excelで一度再計算して保存してから実行してください。"
        )
    raise ValueError(f"循環参照を検出しました: {coordinate}")


def count_paid_leave_half_units(
    sheet: Worksheet, employee: Employee, columns: list[int],
) -> int:
    units = {"年休": 2, "年/": 1, "/年": 1}
    return sum(
        units.get(normalize(sheet.cell(employee.row + 1, col).value), 0)
        for col in columns
    )


def build_holiday_targets(
    formula_sheet: Worksheet,
    cached_sheet: Worksheet,
    employees: list[Employee],
    columns: list[int],
) -> dict[str, HolidayTarget]:
    public_column = find_header_column(formula_sheet, "公休数")
    targets = {}
    for employee in employees:
        cell = formula_sheet.cell(employee.row, public_column)
        public_days = resolve_numeric_cell(
            formula_sheet, cached_sheet, cell.coordinate,
        )
        targets[employee.employee_id] = HolidayTarget(
            public_half_units=numeric_to_half_units(
                public_days, f"{cell.coordinate} 公休数",
            ),
            paid_half_units=count_paid_leave_half_units(
                formula_sheet, employee, columns,
            ),
        )
    return targets


def fixed_off_half_units(work) -> int:
    return {"休": 2, "半/": 1, "/半": 1}.get(work, 0)


def validate_holiday_feasibility(
    sheet: Worksheet,
    employees: list[Employee],
    columns: list[int],
    holiday_targets: dict[str, HolidayTarget],
) -> None:
    errors = []
    for employee in employees:
        fixed_units = flexible_days = 0
        for column in columns:
            work = normalize(sheet.cell(employee.row, column).value)
            reason = normalize(sheet.cell(employee.row + 1, column).value)
            if is_fixed_shift(work, reason):
                fixed_units += fixed_off_half_units(work)
            else:
                flexible_days += 1

        target = holiday_targets[employee.employee_id].total_half_units
        maximum = fixed_units + flexible_days * 2
        if target < fixed_units:
            errors.append(
                f"{employee.name}: 固定済み休日が必要休日数を超えています "
                f"(固定 {fixed_units / 2:g}日 / 必要 {target / 2:g}日)"
            )
        elif target > maximum:
            errors.append(
                f"{employee.name}: 必要休日数をすべて割り当てられません "
                f"(最大 {maximum / 2:g}日 / 必要 {target / 2:g}日)"
            )
        elif (target - fixed_units) % 2:
            # 自動生成は終日勤務・終日休のみ。端数は固定半休で合わせる。
            errors.append(
                f"{employee.name}: 必要休日数と固定半休の組み合わせでは"
                "残りを0日にできません "
                f"(固定 {fixed_units / 2:g}日 / 必要 {target / 2:g}日)"
            )
    if errors:
        raise ValueError(
            "必要休日数を正確に割り当てられません。\n"
            + "\n".join(f" - {error}" for error in errors)
        )


def count_written_off_half_units(
    sheet: Worksheet, employee: Employee, columns: list[int],
) -> int:
    return sum(
        fixed_off_half_units(normalize(sheet.cell(employee.row, col).value))
        for col in columns
    )


def validate_written_holidays(
    sheet: Worksheet,
    employees: list[Employee],
    columns: list[int],
    holiday_targets: dict[str, HolidayTarget],
) -> None:
    errors = []
    for employee in employees:
        actual = count_written_off_half_units(sheet, employee, columns)
        target = holiday_targets[employee.employee_id].total_half_units
        if actual != target:
            errors.append(
                f"{employee.name}: 必要 {target / 2:g}日 / 出力 {actual / 2:g}日"
            )
    if errors:
        raise RuntimeError(
            "出力後の休日数が一致していません。\n"
            + "\n".join(f" - {error}" for error in errors)
        )


def validate_sunday_fixed_staffing(
    sheet: Worksheet, employees: list[Employee], columns: list[int],
) -> None:
    members_by_ward = {
        ward: [employee for employee in employees if employee.ward == ward]
        for ward in sorted({employee.ward for employee in employees})
    }
    errors = []
    for column in columns:
        if weekday(sheet, column) != "日":
            continue
        day = day_label(sheet, column)
        for ward, members in members_by_ward.items():
            fixed_present = flexible = 0
            for employee in members:
                work = normalize(sheet.cell(employee.row, column).value)
                reason = normalize(sheet.cell(employee.row + 1, column).value)
                if not is_fixed_shift(work, reason):
                    flexible += 1
                elif work != "休":
                    fixed_present += 1
            if fixed_present > 2:
                errors.append(
                    f"{day}日(日) {ward}病棟: 固定出勤が{fixed_present}名あり、"
                    "2名勤務にできません。"
                )
            if fixed_present + flexible < 2:
                errors.append(
                    f"{day}日(日) {ward}病棟: 固定勤務の影響で"
                    "2名を確保できません。"
                )
    if errors:
        raise ValueError(
            "日曜日の固定勤務に矛盾があります。\n"
            + "\n".join(f" - {error}" for error in errors)
        )


def capture_formulas(workbook) -> dict[tuple[str, str], str]:
    return {
        (sheet.title, cell.coordinate): cell.value
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.data_type == "f"
    }


def validate_formulas(workbook, formulas: dict[tuple[str, str], str]) -> None:
    for (sheet_name, address), expected in formulas.items():
        if workbook[sheet_name][address].value != expected:
            raise RuntimeError(f"数式が変更されました: {sheet_name}!{address}")


def validate_fixed_shifts(
    sheet: Worksheet, fixed_shifts: dict[CellKey, str],
) -> None:
    for (row, column), expected in fixed_shifts.items():
        cell = sheet.cell(row, column)
        if normalize(cell.value) != expected:
            raise RuntimeError(
                f"固定勤務が変更されました: {cell.coordinate} "
                f"(期待: {expected} / 実際: {normalize(cell.value)})"
            )
