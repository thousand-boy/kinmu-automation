"""職員マスタの検証と、勤務表の職員行との対応付け。"""

from __future__ import annotations

from collections import Counter
import re

from openpyxl.worksheet.worksheet import Worksheet

from .constants import MASTER_HEADERS, MASTER_SHEET_NAME, WARD_STAFF_LIMITS
from .models import Employee, EmployeeMaster


def normalize(value):
    if isinstance(value, str):
        return value.strip() or None
    return value


def employee_name(raw_value) -> str | None:
    """表示名の役割表記（専従など）を除いて照合する。"""
    value = normalize(raw_value)
    if value is None:
        return None
    return re.split(r"[（(]", str(value), maxsplit=1)[0].strip()


def parse_flag(value, coordinate: str) -> bool:
    value = normalize(value)
    if value in {True, "1", "○", "〇", "true", "TRUE", "yes", "YES", "有", "あり"}:
        return True
    if value in {
        None, False, "0", "×", "✕", "false", "FALSE", "no", "NO", "無", "なし",
    }:
        return False
    raise ValueError(
        f"{MASTER_SHEET_NAME}!{coordinate}: "
        f"「{value}」を○/×として認識できません。"
    )


def parse_ward(value, coordinate: str) -> int:
    value = normalize(value)
    error = (
        f"{MASTER_SHEET_NAME}!{coordinate}: 病棟は整数で入力してください。"
    )
    if isinstance(value, bool):
        raise ValueError(error)
    if isinstance(value, (int, float)):
        if not float(value).is_integer():
            raise ValueError(error)
        ward = int(value)
    else:
        text = str(value or "")
        if not re.fullmatch(r"\d+", text):
            raise ValueError(error)
        ward = int(text)

    if ward not in WARD_STAFF_LIMITS:
        supported = ", ".join(map(str, sorted(WARD_STAFF_LIMITS)))
        raise ValueError(
            f"{MASTER_SHEET_NAME}!{coordinate}: "
            f"病棟「{ward}」の勤務条件が未定義です。 対応病棟: {supported}"
        )
    return ward


def load_employee_master(sheet: Worksheet) -> list[EmployeeMaster]:
    header_columns = {}
    for cell in sheet[1]:
        header = normalize(cell.value)
        if header not in MASTER_HEADERS:
            continue
        if header in header_columns:
            raise ValueError(
                f"{MASTER_SHEET_NAME}: 見出し「{header}」が重複しています。"
            )
        header_columns[header] = cell.column

    missing = [name for name in MASTER_HEADERS if name not in header_columns]
    if missing:
        raise ValueError(
            f"{MASTER_SHEET_NAME}: 必須見出しがありません: {', '.join(missing)}"
        )

    records = []
    employee_ids = set()
    active_names = set()
    for row in range(2, sheet.max_row + 1):
        cells = {
            name: sheet.cell(row, column)
            for name, column in header_columns.items()
        }
        values = {name: normalize(cell.value) for name, cell in cells.items()}
        if all(value is None for value in values.values()):
            continue
        for name in ("職員ID", "表示名", "病棟"):
            if values[name] is None:
                raise ValueError(
                    f"{MASTER_SHEET_NAME}: {row}行目の{name}が空欄です。"
                )

        employee_id = str(values["職員ID"])
        if employee_id in employee_ids:
            raise ValueError(
                f"{MASTER_SHEET_NAME}: 職員ID「{employee_id}」が重複しています。"
            )
        employee_ids.add(employee_id)

        record = EmployeeMaster(
            employee_id=employee_id,
            display_name=str(values["表示名"]),
            ward=parse_ward(values["病棟"], cells["病棟"].coordinate),
            is_dedicated=parse_flag(values["専従"], cells["専従"].coordinate),
            is_assigned=parse_flag(values["専任"], cells["専任"].coordinate),
            is_ot=parse_flag(values["作業療法士"], cells["作業療法士"].coordinate),
            is_active=parse_flag(values["在籍"], cells["在籍"].coordinate),
        )
        if record.is_dedicated and record.is_assigned:
            raise ValueError(
                f"{MASTER_SHEET_NAME}: {row}行目で"
                "「専従」と「専任」が両方○になっています。"
            )
        if record.is_active and record.display_name in active_names:
            raise ValueError(
                f"{MASTER_SHEET_NAME}: 在籍職員の表示名"
                f"「{record.display_name}」が重複しています。"
            )
        if record.is_active:
            active_names.add(record.display_name)
        records.append(record)

    if not records:
        raise ValueError(f"{MASTER_SHEET_NAME}: 職員情報がありません。")
    if not active_names:
        raise ValueError(f"{MASTER_SHEET_NAME}: 在籍○の職員がいません。")
    return records


def employee_display_label(record: EmployeeMaster) -> str:
    if record.is_dedicated:
        role = "専従"
    elif record.is_assigned:
        role = "専任"
    elif record.is_ot:
        role = "作業療法士"
    else:
        return record.display_name
    return f"{record.display_name}（{role}）"


def _schedule_aliases(
    records: list[EmployeeMaster],
) -> dict[str, EmployeeMaster]:
    aliases = {}
    for record in records:
        for alias in {record.employee_id, record.display_name}:
            existing = aliases.get(alias)
            if existing and existing.employee_id != record.employee_id:
                raise ValueError(
                    f"{MASTER_SHEET_NAME}: 「{alias}」だけでは"
                    "職員を一意に判定できません。"
                )
            aliases[alias] = record
    return aliases


def _staff_rows(
    sheet: Worksheet, master_records: list[EmployeeMaster],
) -> dict[str, int]:
    aliases = _schedule_aliases(master_records)
    rows = {}
    for row in range(1, sheet.max_row + 1):
        if normalize(sheet.cell(row, 2).value) != "勤務希望":
            continue
        label = employee_name(sheet.cell(row, 1).value)
        if label is None:
            continue
        record = aliases.get(label)
        if record is None:
            raise ValueError(
                f"{sheet.title}!A{row}: 「{label}」が"
                f"{MASTER_SHEET_NAME}に見つかりません。"
            )
        if record.employee_id in rows:
            raise ValueError(
                f"職員「{record.display_name}」の勤務希望行が複数あります。"
            )
        rows[record.employee_id] = row
    return rows


def discover_employees(
    sheet: Worksheet, master_records: list[EmployeeMaster],
) -> list[Employee]:
    """役割・病棟はマスタを採用し、在籍者だけを返す。"""
    rows = _staff_rows(sheet, master_records)
    active = [record for record in master_records if record.is_active]
    missing = [r.display_name for r in active if r.employee_id not in rows]
    if missing:
        raise ValueError(
            f"勤務表に在籍職員の行が見つかりません: {', '.join(missing)}"
        )
    return [
        Employee(
            employee_id=record.employee_id,
            name=record.display_name,
            row=rows[record.employee_id],
            ward=record.ward,
            is_dedicated=record.is_dedicated,
            is_assigned=record.is_assigned,
            is_ot=record.is_ot,
        )
        for record in active
    ]


def validate_employee_structure(employees: list[Employee]) -> None:
    counts = Counter(employee.ward for employee in employees)
    missing = WARD_STAFF_LIMITS.keys() - counts.keys()
    if missing:
        raise ValueError(
            "在籍職員が登録されていない病棟があります: "
            + ", ".join(map(str, sorted(missing)))
        )
    for ward in sorted(WARD_STAFF_LIMITS):
        if counts[ward] < 2:
            raise ValueError(
                f"{ward}病棟の在籍職員が{counts[ward]}名しかいません。"
                "日曜日を2名勤務にできません。"
            )
    if not any(employee.is_ot for employee in employees):
        raise ValueError("在籍職員に作業療法士が登録されていません。")
