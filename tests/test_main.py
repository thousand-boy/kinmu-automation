from dataclasses import replace

import pytest
from openpyxl import Workbook
from openpyxl.formula import Tokenizer
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from kinmu.models import Employee, EmployeeMaster, HolidayTarget
from kinmu.optimizer import solve_schedule
from kinmu.reporting import four_day_violations
from kinmu.roster import (
    discover_employees,
    employee_name,
    load_employee_master,
    normalize,
    parse_flag,
    parse_ward,
)
from kinmu.validation import (
    build_holiday_targets,
    is_fixed_shift,
    numeric_to_half_units,
    resolve_numeric_cell,
    validate_holiday_feasibility,
)
from kinmu.workbook import (
    synchronize_employee_rows,
    update_holiday_summary_formulas,
)


def create_master_sheet():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '職員マスタ'
    sheet.append(['職員ID', '表示名', '病棟', '専従', '専任', '作業療法士', '在籍'])
    sheet.append(['A', 'A', 1, '○', None, None, '○'])
    sheet.append(['B', 'B', 1, None, '○', None, '○'])
    sheet.append(['D', 'D', 1, None, None, '○', '○'])
    return sheet


def create_schedule_sheet(works=None, reasons=None, *, public_days=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Sheet1'
    for column, (date, day) in enumerate(zip(range(1, 5), '月火水木'), start=3):
        sheet.cell(2, column, date)
        sheet.cell(3, column, day)
    sheet['A4'] = 'A'
    sheet['B4'] = '勤務希望'
    sheet['B5'] = '理由欄'
    if works is None:
        works = [None] * 4
    if reasons is None:
        reasons = [None] * 4
    for column, value in enumerate(works, start=3):
        sheet.cell(row=4, column=column, value=value)
    for column, value in enumerate(reasons, start=3):
        sheet.cell(row=5, column=column, value=value)
    if public_days is not None:
        sheet['H1'] = '公休数'
        sheet['H4'] = public_days
    return sheet


def create_employee():
    return Employee(
        employee_id='A', name='A', row=4, ward=1,
        is_dedicated=True, is_assigned=False, is_ot=True,
    )


def target(days: float):
    return {
        'A': HolidayTarget(public_half_units=int(days * 2), paid_half_units=0)
    }


def test_normalize():
    assert normalize(' 休 ') == '休'
    assert normalize('') is None
    assert normalize('   ') is None
    assert normalize(None) is None
    assert normalize(10) == 10


@pytest.mark.parametrize(
    ('raw_value', 'expected'),
    [
        ('A（専従）', 'A'),
        ('B（専任）', 'B'),
        ('D（作業療法士）', 'D'),
        ('A(専従)', 'A'),
        (' C ', 'C'),
        (None, None),
    ],
)
def test_employee_name(raw_value, expected):
    assert employee_name(raw_value) == expected


@pytest.mark.parametrize(
    ('value', 'expected'),
    [
        ('○', True), ('〇', True), (1, True), (True, True),
        (None, False), ('×', False), (0, False), (False, False),
    ],
)
def test_parse_flag(value, expected):
    assert parse_flag(value, 'A1') is expected


def test_parse_flag_rejects_unknown_value():
    with pytest.raises(ValueError, match='認識できません'):
        parse_flag('△', 'A1')


@pytest.mark.parametrize(
    ('value', 'expected'), [(1, 1), (2, 2), ('3', 3), (1.0, 1)],
)
def test_parse_ward(value, expected):
    assert parse_ward(value, 'C2') == expected


@pytest.mark.parametrize('value', [1.5, 'A', 4, None])
def test_parse_ward_rejects_invalid_value(value):
    with pytest.raises(ValueError):
        parse_ward(value, 'C2')


def test_load_employee_master():
    records = load_employee_master(create_master_sheet())
    assert len(records) == 3
    assert records[0] == EmployeeMaster(
        employee_id='A', display_name='A', ward=1,
        is_dedicated=True, is_assigned=False, is_ot=False, is_active=True,
    )
    assert records[1].is_assigned is True
    assert records[2].is_ot is True


def test_load_employee_master_rejects_duplicate_id():
    sheet = create_master_sheet()
    sheet.append(['A', '別のA', 1, None, None, None, '○'])
    with pytest.raises(ValueError, match='職員ID.*重複'):
        load_employee_master(sheet)


def test_load_employee_master_rejects_dedicated_and_assigned():
    sheet = create_master_sheet()
    sheet.append(['X', 'X', 1, '○', '○', None, '○'])
    with pytest.raises(ValueError, match='専従.*専任'):
        load_employee_master(sheet)


def test_discover_employees_uses_master_as_source_of_truth():
    workbook = Workbook()
    schedule = workbook.active
    schedule.title = 'Sheet1'
    schedule['A4'] = 'A（専任）'
    schedule['B4'] = '勤務希望'
    master_records = [
        EmployeeMaster(
            employee_id='A', display_name='A', ward=2,
            is_dedicated=False, is_assigned=False, is_ot=True, is_active=True,
        )
    ]
    employees = discover_employees(schedule, master_records)
    employee = employees[0]
    assert employee.employee_id == 'A'
    assert employee.ward == 2
    assert employee.is_assigned is False
    assert employee.is_ot is True


def test_discover_employees_rejects_missing_employee():
    workbook = Workbook()
    schedule = workbook.active
    schedule.title = 'Sheet1'
    master_records = [
        EmployeeMaster(
            employee_id='A', display_name='A', ward=1,
            is_dedicated=True, is_assigned=False, is_ot=False, is_active=True,
        )
    ]
    with pytest.raises(ValueError, match='勤務表に在籍職員'):
        discover_employees(schedule, master_records)


def test_synchronize_employee_rows_adds_pair_and_preserves_features():
    workbook = Workbook()
    schedule = workbook.active
    schedule.title = 'Sheet1'
    schedule['A4'] = 'A'
    schedule['B4'] = '勤務希望'
    schedule['B5'] = '理由欄'
    schedule.merge_cells('A4:A5')
    schedule['AI4'] = '=AS8'
    schedule['AJ4'] = '=COUNTIF(C5:G5,"年休")'
    schedule['AK4'] = '=COUNTIF(C4:G4,"休")'
    schedule['AL4'] = '=AI4+AJ4-AK4'
    schedule['B6'] = '階休み'
    schedule['C6'] = '=COUNTIF(C4:C5,"休")'
    schedule.row_dimensions[4].height = 27
    schedule.row_dimensions[5].height = 19
    schedule['C4'].fill = PatternFill('solid', fgColor='FFF2CC')
    work_validation = DataValidation(type='list', formula1='$AN$8:$AN$9')
    reason_validation = DataValidation(type='list', formula1='"希望,年休"')
    schedule.add_data_validation(work_validation)
    schedule.add_data_validation(reason_validation)
    work_validation.add('C4:G4')
    reason_validation.add('C5:G5')
    schedule.conditional_formatting.add(
        'C4:G5',
        FormulaRule(
            formula=['C4="休"'],
            fill=PatternFill('solid', fgColor='DDDDDD'),
        ),
    )
    master_records = [
        EmployeeMaster('A', 'A', 1, True, False, False, True),
        EmployeeMaster('Q', 'Q', 1, False, False, True, True),
    ]
    added = synchronize_employee_rows(schedule, master_records)
    assert added == ['Q']
    assert schedule['A6'].value == 'Q（作業療法士）'
    assert schedule['B6'].value == '勤務希望'
    assert schedule['B7'].value == '理由欄'
    assert 'A6:A7' in {str(item) for item in schedule.merged_cells.ranges}
    assert schedule.row_dimensions[6].height == 27
    assert schedule.row_dimensions[7].height == 19
    assert schedule['C6'].fill.fgColor.rgb == schedule['C4'].fill.fgColor.rgb
    # 共有の公休日数も挿入位置より下にあるため、参照先を移動する。
    assert schedule['AI6'].value == '=AS10'
    assert schedule['AJ6'].value == '=COUNTIF(C7:G7,"年休")'
    assert schedule['AK6'].value == '=COUNTIF(C6:G6,"休")'
    assert schedule['AL6'].value == '=AI6+AJ6-AK6'
    assert schedule['C8'].value == '=COUNTIF(C4:C7,"休")'
    # 病棟集計は広げるが、直前の職員の個人集計には新しい行を含めない。
    assert schedule['AJ4'].value == '=COUNTIF(C5:G5,"年休")'
    assert schedule['AK4'].value == '=COUNTIF(C4:G4,"休")'
    validation_ranges = ' '.join(
        str(validation.sqref)
        for validation in schedule.data_validations.dataValidation
    )
    assert 'C6:G6' in validation_ranges
    assert 'C7:G7' in validation_ranges
    assert work_validation.formula1 == '$AN$10:$AN$11'
    conditional_ranges = ' '.join(
        str(item.sqref) for item in schedule.conditional_formatting
    )
    assert 'C6:G7' in conditional_ranges
    assert any(
        rule.formula == ['C6="休"']
        for rules in schedule.conditional_formatting._cf_rules.values()
        for rule in rules
    )
    employees = discover_employees(schedule, master_records)
    assert [(employee.name, employee.row) for employee in employees] == [
        ('A', 4), ('Q', 6),
    ]


def test_synchronize_employee_rows_keeps_inactive_rows_but_excludes_them():
    workbook = Workbook()
    schedule = workbook.active
    schedule.title = 'Sheet1'
    schedule['A4'] = 'A'
    schedule['B4'] = '勤務希望'
    schedule['B5'] = '理由欄'
    schedule['A6'] = 'X'
    schedule['B6'] = '勤務希望'
    schedule['B7'] = '理由欄'
    schedule['B8'] = '階休み'
    master_records = [
        EmployeeMaster('A', 'A', 1, True, False, False, True),
        EmployeeMaster('X', 'X', 1, False, False, False, False),
    ]
    assert synchronize_employee_rows(schedule, master_records) == []
    assert schedule['A6'].value == 'X'
    employees = discover_employees(schedule, master_records)
    assert [employee.name for employee in employees] == ['A']


def test_synchronize_employee_rows_moves_dimensions_below_insertion():
    schedule = create_schedule_sheet()
    schedule['B6'] = '階休み'
    schedule.row_dimensions[4].height = 27
    schedule.row_dimensions[5].height = 19
    schedule.row_dimensions[6].height = 32
    schedule.row_dimensions[7].height = 45
    schedule.row_dimensions[7].hidden = True
    schedule.row_dimensions[7].outlineLevel = 2
    records = [
        EmployeeMaster('A', 'A', 1, True, False, False, True),
        EmployeeMaster('Q', 'Q', 1, False, False, True, True),
    ]

    synchronize_employee_rows(schedule, records)

    assert schedule.row_dimensions[6].height == 27
    assert schedule.row_dimensions[7].height == 19
    assert schedule.row_dimensions[7].hidden is False
    assert schedule.row_dimensions[8].height == 32
    assert schedule.row_dimensions[9].height == 45
    assert schedule.row_dimensions[9].hidden is True
    assert schedule.row_dimensions[9].outlineLevel == 2
    assert schedule.row_dimensions[9].index == 9


def test_synchronize_employee_rows_shifts_cross_sheet_references_only():
    schedule = create_schedule_sheet()
    schedule['B6'] = '階休み'
    other = schedule.parent.create_sheet('Summary')
    other['A1'] = '=Sheet1!C14'
    other['A2'] = "='Sheet1'!$C$14:$D$15"
    other['A3'] = '=C14+Sheet1!C14'
    other['A4'] = '=SUM(C14:D15)'
    other['A5'] = '=IF(Sheet1!C14="C14",1,0)'
    records = [
        EmployeeMaster('A', 'A', 1, True, False, False, True),
        EmployeeMaster('Q', 'Q', 1, False, False, True, True),
    ]

    synchronize_employee_rows(schedule, records)

    assert other['A1'].value == '=Sheet1!C16'
    assert other['A2'].value == "='Sheet1'!$C$16:$D$17"
    assert other['A3'].value == '=C14+Sheet1!C16'
    assert other['A4'].value == '=SUM(C14:D15)'
    assert other['A5'].value == '=IF(Sheet1!C16="C14",1,0)'


def create_summary_sheet():
    sheet = Workbook().active
    sheet.title = 'Sheet1'
    for row, label in ((4, 'A'), (6, 'B'), (10, 'C'), (16, 'D')):
        sheet.cell(row, 1).value = label
        sheet.cell(row, 2).value = '勤務希望'
        sheet.cell(row + 1, 2).value = '理由欄'
    sheet['C4'] = '休'
    sheet['C5'] = '希望'
    sheet['C6'] = '休'
    sheet['C7'] = '年休'
    sheet['C10'] = '半/'
    sheet['D10'] = '/半'
    for row, label in ((8, '階休み'), (14, '階休み'), (18, '階休み'), (20, '全体休み')):
        sheet.cell(row, 2).value = label
        sheet.cell(row, 3).value = '=99'
    sheet['E8'] = '対象日以外は維持'
    employees = [
        replace(create_employee(), employee_id='A', name='A', row=4, ward=1),
        replace(create_employee(), employee_id='C', name='C', row=10, ward=2),
        replace(create_employee(), employee_id='D', name='D', row=16, ward=3),
    ]
    return sheet, employees


def formula_references(formula):
    return {
        item.value for item in Tokenizer(formula).items
        if item.type == 'OPERAND' and item.subtype == 'RANGE'
    }


def test_holiday_summaries_exclude_inactive_and_preserve_personal_cells():
    sheet, employees = create_summary_sheet()
    before = {
        (row, col): (
            sheet.cell(row, col).value, str(sheet.cell(row, col)._style),
        )
        for row in (4, 5, 6, 7, 10, 11, 16, 17)
        for col in range(1, 39)
    }
    update_holiday_summary_formulas(sheet, employees, [3, 4])
    assert formula_references(sheet['C8'].value) == {'C4'}
    assert formula_references(sheet['C14'].value) == {'C10'}
    assert formula_references(sheet['C18'].value) == {'C16'}
    assert formula_references(sheet['C20'].value) == {'C8', 'C14', 'C18'}
    assert formula_references(sheet['D14'].value) == {'D10'}
    assert before == {
        (row, col): (
            sheet.cell(row, col).value, str(sheet.cell(row, col)._style),
        )
        for row in (4, 5, 6, 7, 10, 11, 16, 17)
        for col in range(1, 39)
    }
    assert sheet['E8'].value == '対象日以外は維持'


def test_holiday_summaries_include_reactivated_employee():
    sheet, employees = create_summary_sheet()
    update_holiday_summary_formulas(sheet, employees, [3])
    assert 'C6' not in formula_references(sheet['C8'].value)
    employees.append(
        replace(create_employee(), employee_id='B', name='B', row=6, ward=1)
    )
    update_holiday_summary_formulas(sheet, employees, [3])
    assert formula_references(sheet['C8'].value) == {'C4', 'C6'}
    assert sheet['C6'].value == '休'
    assert sheet['C7'].value == '年休'


def test_holiday_summaries_use_master_ward_not_physical_row_group():
    sheet, employees = create_summary_sheet()
    employees[0] = replace(employees[0], ward=2)
    update_holiday_summary_formulas(sheet, employees, [3])
    assert sheet['C8'].value == '=0'
    assert formula_references(sheet['C14'].value) == {'C4', 'C10'}


@pytest.mark.parametrize('missing_row', [8, 14, 20])
def test_holiday_summaries_reject_missing_layout_without_partial_changes(
    missing_row,
):
    sheet, employees = create_summary_sheet()
    sheet.cell(missing_row, 2).value = None
    with pytest.raises(ValueError, match='休み集計行'):
        update_holiday_summary_formulas(sheet, employees, [3])
    assert sheet['C8'].value == '=99'
    assert sheet['C14'].value == '=99'
    assert sheet['C18'].value == '=99'
    assert sheet['C20'].value == '=99'


@pytest.mark.parametrize(
    ('work', 'reason', 'expected'),
    [
        ('休', '希望', True), ('出', '委員会', True),
        ('半/', '年/', True), ('/半', '/年', True),
        ('休', None, False), ('半/', None, False),
        (None, '希望', False), (None, None, False),
    ],
)
def test_is_fixed_shift(work, reason, expected):
    assert is_fixed_shift(work, reason) is expected


@pytest.mark.parametrize(
    ('value', 'expected'), [(10, 20), (1, 2), (0.5, 1), (1.5, 3), (0, 0)],
)
def test_numeric_to_half_units(value, expected):
    assert numeric_to_half_units(value, 'test') == expected


def test_numeric_to_half_units_rejects_quarter_day():
    with pytest.raises(ValueError, match='0.5日単位'):
        numeric_to_half_units(1.25, 'test')


@pytest.mark.parametrize(
    'value', [True, False, -0.5, float('inf'), float('nan')],
)
def test_numeric_to_half_units_rejects_invalid_numbers(value):
    with pytest.raises(ValueError, match='有限な日数'):
        numeric_to_half_units(value, '公休数')


def test_numeric_to_half_units_rejects_overflow_when_scaling():
    with pytest.raises(ValueError, match='大きすぎます'):
        numeric_to_half_units(1e308, '公休数')


def test_resolve_numeric_cell_follows_absolute_reference_chain():
    sheet = Workbook().active
    sheet['A1'] = '=$B$2'
    sheet['B2'] = '=C3'
    sheet['C3'] = 1.5
    assert resolve_numeric_cell(sheet, sheet, 'A1') == 1.5


def test_resolve_numeric_cell_uses_cached_formula_value():
    sheet = Workbook().active
    cache = Workbook().active
    sheet['A1'] = '=SUM(B1:B2)'
    cache['A1'] = 2.5
    assert resolve_numeric_cell(sheet, cache, 'A1') == 2.5


def test_resolve_numeric_cell_rejects_cycles():
    sheet = Workbook().active
    sheet['A1'] = '=B1'
    sheet['B1'] = '=A1'
    with pytest.raises(ValueError, match='循環参照'):
        resolve_numeric_cell(sheet, sheet, 'A1')


@pytest.mark.parametrize('value', [True, False, None, '=SUM(B1:B2)'])
def test_resolve_numeric_cell_requires_numeric_value_or_cache(value):
    sheet = Workbook().active
    sheet['A1'] = value
    with pytest.raises(ValueError, match='数値を取得できません'):
        resolve_numeric_cell(sheet, sheet, 'A1')


def test_build_holiday_targets_adds_paid_leave():
    sheet = create_schedule_sheet(
        works=['半/', None, None, None],
        reasons=['年/', None, None, None],
        public_days=1.5,
    )
    employee = create_employee()
    targets = build_holiday_targets(sheet, sheet, [employee], [3, 4, 5, 6])
    assert targets['A'].public_half_units == 3
    assert targets['A'].paid_half_units == 1
    assert targets['A'].total_half_units == 4


def test_holiday_feasibility_rejects_wrong_half_day_parity():
    sheet = create_schedule_sheet()
    employee = create_employee()
    with pytest.raises(ValueError, match='残りを0'):
        validate_holiday_feasibility(
            sheet, [employee], [3, 4, 5, 6],
            {'A': HolidayTarget(public_half_units=1, paid_half_units=0)},
        )


def test_holiday_feasibility_accepts_fixed_half_day():
    sheet = create_schedule_sheet(
        works=['半/', None, None, None], reasons=['年/', None, None, None],
    )
    employee = create_employee()
    validate_holiday_feasibility(
        sheet, [employee], [3, 4, 5, 6],
        {'A': HolidayTarget(public_half_units=0, paid_half_units=1)},
    )


def test_strict_mode_allocates_exact_holidays_without_four_consecutive():
    sheet = create_schedule_sheet()
    employee = create_employee()
    result = solve_schedule(sheet, [employee], [3, 4, 5, 6], target(1))
    assert result.mode == 'strict'
    assert result.solver.value(result.data.holiday_off['A']) == 2
    assert four_day_violations(result.solver, result.data) == []


def test_fallback_relaxes_four_consecutive_but_keeps_holiday_target():
    sheet = create_schedule_sheet()
    employee = create_employee()
    # 休日0日なら4日すべて出勤が必要になり、通常モードでは解けない。
    result = solve_schedule(sheet, [employee], [3, 4, 5, 6], target(0))
    assert result.mode == 'relaxed'
    assert result.solver.value(result.data.holiday_off['A']) == 0
    violations = four_day_violations(result.solver, result.data)
    assert len(violations) == 1
    assert violations[0][1] is False


def test_fixed_four_consecutive_is_kept_and_warned():
    sheet = create_schedule_sheet(works=['出'] * 4, reasons=['希望'] * 4)
    employee = create_employee()
    result = solve_schedule(sheet, [employee], [3, 4, 5, 6], target(0))
    assert result.mode == 'strict'
    violations = four_day_violations(result.solver, result.data)
    assert len(violations) == 1
    assert violations[0][1] is True


def create_sunday_schedule():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Sheet1'
    sheet['C2'] = 1
    sheet['C3'] = '日'
    rows = [(4, 'A'), (6, 'B'), (8, 'C')]
    for row, name in rows:
        sheet.cell(row=row, column=1, value=name)
        sheet.cell(row=row, column=2, value='勤務希望')
        sheet.cell(row=row + 1, column=2, value='理由欄')
    return sheet


def create_sunday_employees():
    return [
        create_employee(),
        replace(
            create_employee(), employee_id='B', name='B', row=6,
            is_dedicated=False, is_assigned=True, is_ot=False,
        ),
        replace(
            create_employee(), employee_id='C', name='C', row=8,
            is_dedicated=False, is_assigned=False, is_ot=False,
        ),
    ]


def test_sunday_has_exactly_two_workers():
    sheet = create_sunday_schedule()
    employees = create_sunday_employees()
    targets = {
        'A': HolidayTarget(public_half_units=0, paid_half_units=0),
        'B': HolidayTarget(public_half_units=0, paid_half_units=0),
        'C': HolidayTarget(public_half_units=2, paid_half_units=0),
    }
    result = solve_schedule(sheet, employees, [3], targets)
    working = sum(
        result.solver.value(result.data.full_work[employee.row, 3])
        for employee in employees
    )
    assert working == 2


def test_sunday_three_fixed_workers_is_infeasible():
    sheet = create_sunday_schedule()
    employees = create_sunday_employees()
    for employee in employees:
        sheet.cell(row=employee.row, column=3, value='出')
        sheet.cell(row=employee.row + 1, column=3, value='希望')
    targets = {
        employee.employee_id: HolidayTarget(0, 0)
        for employee in employees
    }
    with pytest.raises(RuntimeError, match='日曜日各病棟2名'):
        solve_schedule(sheet, employees, [3], targets)
