import shutil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter

from kinmu.application import run, save_safely
from kinmu.constants import (
    MASTER_SHEET_NAME, SCHEDULE_SHEET_NAME, WARD_STAFF_LIMITS,
)
from kinmu.optimizer import solve_schedule
from kinmu.reporting import four_day_violations
from kinmu.roster import (
    discover_employees,
    employee_name,
    load_employee_master,
    normalize,
    validate_employee_structure,
)
from kinmu.validation import (
    build_holiday_targets,
    capture_formulas,
    count_written_off_half_units,
    is_fixed_shift,
    validate_fixed_shifts,
    validate_formulas,
    validate_holiday_feasibility,
    validate_input,
    validate_sunday_fixed_staffing,
    validate_written_holidays,
)
from kinmu.workbook import (
    day_columns,
    day_label,
    update_holiday_summary_formulas,
    weekday,
    write_schedule,
)

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SAMPLE_FILE = PROJECT_ROOT / 'sample' / 'kinmu_sample.xlsx'


def assert_active_holiday_summary_references(sheet, employees, columns):
    summary_rows = [
        row for row in range(1, sheet.max_row + 1)
        if sheet.cell(row, 2).value == '階休み'
    ]
    assert len(summary_rows) == len(WARD_STAFF_LIMITS)
    for ward, summary_row in zip(sorted(WARD_STAFF_LIMITS), summary_rows):
        for col in columns:
            references = formula_references(sheet.cell(summary_row, col).value)
            expected = {
                f'{get_column_letter(col)}{employee.row}'
                for employee in employees if employee.ward == ward
            }
            assert references == expected
    overall_row = next(
        row for row in range(1, sheet.max_row + 1)
        if sheet.cell(row, 2).value == '全体休み'
    )
    for col in columns:
        formula = sheet.cell(overall_row, col).value
        references = formula_references(formula)
        assert references == {
            f'{get_column_letter(col)}{row}' for row in summary_rows
        }


def formula_references(formula):
    return {
        token.value for token in Tokenizer(formula).items
        if token.type == 'OPERAND' and token.subtype == 'RANGE'
    }


def output_presence(sheet, employee, column):
    """半休も出勤人数に含める（休日数の集計とは異なる）。"""
    return int(normalize(sheet.cell(employee.row, column).value) != '休')


def test_real_sample_end_to_end(tmp_path):
    """公開サンプルの読み込み、最適化、保存、再読み込みを検証する。"""
    assert SAMPLE_FILE.exists(), 'sample/kinmu_sample.xlsx が見つかりません。'
    test_input_dir = tmp_path / 'input'
    test_output_dir = tmp_path / 'output'
    test_input_dir.mkdir()
    test_output_dir.mkdir()
    test_input_file = test_input_dir / 'kinmu_sample.xlsx'
    test_output_file = test_output_dir / 'kinmu_output.xlsx'
    shutil.copy2(SAMPLE_FILE, test_input_file)
    original_input = test_input_file.read_bytes()
    workbook = load_workbook(test_input_file, data_only=False)
    cached_workbook = load_workbook(test_input_file, data_only=True)
    try:
        assert SCHEDULE_SHEET_NAME in workbook.sheetnames
        assert MASTER_SHEET_NAME in workbook.sheetnames
        schedule_sheet = workbook[SCHEDULE_SHEET_NAME]
        master_sheet = workbook[MASTER_SHEET_NAME]
        cached_sheet = cached_workbook[SCHEDULE_SHEET_NAME]
        master_records = load_employee_master(master_sheet)
        employees = discover_employees(schedule_sheet, master_records)
        validate_employee_structure(employees)
        assert len(employees) == 16
        assert {employee.name for employee in employees} == set(
            'ABCDEFGHIJKLMNOP'
        )
        columns = day_columns(schedule_sheet)
        assert columns
        input_errors = validate_input(schedule_sheet, employees, columns)
        assert input_errors == [], '\n'.join(input_errors)
        original_fixed_shifts = {}
        for employee in employees:
            for column in columns:
                work = normalize(
                    schedule_sheet.cell(employee.row, column).value
                )
                reason = normalize(
                    schedule_sheet.cell(employee.row + 1, column).value
                )
                if is_fixed_shift(work, reason):
                    original_fixed_shifts[employee.row, column] = work
        holiday_targets = build_holiday_targets(
            schedule_sheet, cached_sheet, employees, columns,
        )
        assert len(holiday_targets) == len(employees)
        validate_holiday_feasibility(
            schedule_sheet, employees, columns, holiday_targets,
        )
        validate_sunday_fixed_staffing(schedule_sheet, employees, columns)
        update_holiday_summary_formulas(schedule_sheet, employees, columns)
        assert_active_holiday_summary_references(
            schedule_sheet, employees, columns,
        )
        formulas = capture_formulas(workbook)
        assert formulas
        result = solve_schedule(
            schedule_sheet, employees, columns, holiday_targets,
        )
        assert result.mode in {'strict', 'relaxed'}
        write_schedule(
            schedule_sheet, employees, columns, result.solver, result.data,
        )
        for employee in employees:
            actual = count_written_off_half_units(
                schedule_sheet, employee, columns,
            )
            expected = holiday_targets[employee.employee_id].total_half_units
            assert actual == expected, (
                f'{employee.name}: 必要休日数={expected / 2:g}日 / '
                f'割当={actual / 2:g}日'
            )
        wards = sorted({employee.ward for employee in employees})
        for column in columns:
            if weekday(schedule_sheet, column) != '日':
                continue
            for ward in wards:
                members = [
                    employee for employee in employees if employee.ward == ward
                ]
                working_count = sum(
                    output_presence(schedule_sheet, employee, column)
                    for employee in members
                )
                assert working_count == 2, (
                    f'{day_label(schedule_sheet, column)}日(日) '
                    f'{ward}病棟: 出勤人数={working_count}名'
                )
        for (row, column), expected in original_fixed_shifts.items():
            actual = normalize(schedule_sheet.cell(row, column).value)
            assert actual == expected
        validate_formulas(workbook, formulas)
        four_day = four_day_violations(result.solver, result.data)
        if result.mode == 'strict':
            auto_four_day = [
                warning for warning, fixed_only in four_day if not fixed_only
            ]
            assert auto_four_day == []
        save_safely(
            workbook, formulas, result.data.fixed_shifts,
            employees, columns, holiday_targets, test_output_file,
        )
    finally:
        cached_workbook.close()
        workbook.close()
    assert test_output_file.exists()
    assert test_input_file.read_bytes() == original_input
    output_book = load_workbook(test_output_file, data_only=False)
    try:
        assert SCHEDULE_SHEET_NAME in output_book.sheetnames
        assert MASTER_SHEET_NAME in output_book.sheetnames
        validate_formulas(output_book, formulas)
        output_sheet = output_book[SCHEDULE_SHEET_NAME]
        validate_fixed_shifts(output_sheet, original_fixed_shifts)
        validate_written_holidays(
            output_sheet, employees, columns, holiday_targets,
        )
        assert_active_holiday_summary_references(
            output_sheet, employees, columns,
        )
    finally:
        output_book.close()


def test_real_sample_adds_active_employee_and_keeps_inactive_row(tmp_path):
    """Qを自動追加し、在籍×のPは削除せず最適化対象外にする。"""
    test_input_dir = tmp_path / 'input'
    test_output_dir = tmp_path / 'output'
    test_input_dir.mkdir()
    test_output_dir.mkdir()
    test_input_file = test_input_dir / 'kinmu_sample.xlsx'
    test_output_file = test_output_dir / 'kinmu_output.xlsx'
    shutil.copy2(SAMPLE_FILE, test_input_file)
    workbook = load_workbook(test_input_file, data_only=False)
    try:
        original_schedule = workbook[SCHEDULE_SHEET_NAME]
        original_p_row = next(
            row for row in range(1, original_schedule.max_row + 1)
            if original_schedule.cell(row, 2).value == '勤務希望'
            and employee_name(original_schedule.cell(row, 1).value) == 'P'
        )
        original_p_pair = [
            [original_schedule.cell(row, column).value
             for column in range(1, original_schedule.max_column + 1)]
            for row in (original_p_row, original_p_row + 1)
        ]
        master = workbook[MASTER_SHEET_NAME]
        headers = {normalize(cell.value): cell.column for cell in master[1]}
        for row in range(2, master.max_row + 1):
            if master.cell(row=row, column=headers['職員ID']).value == 'P':
                master.cell(row=row, column=headers['在籍']).value = '×'
                break
        else:
            raise AssertionError('職員マスタにPが見つかりません。')
        master.append(['Q', 'Q', 3, None, None, '○', '○'])
        workbook.save(test_input_file)
    finally:
        workbook.close()
    original_input = test_input_file.read_bytes()
    run(test_input_file, test_output_file)
    assert test_input_file.read_bytes() == original_input
    output_book = load_workbook(test_output_file, data_only=False)
    try:
        schedule = output_book[SCHEDULE_SHEET_NAME]
        master_records = load_employee_master(output_book[MASTER_SHEET_NAME])
        employees = discover_employees(schedule, master_records)
        assert {employee.name for employee in employees} == (
            set('ABCDEFGHIJKLMNO') | {'Q'}
        )
        q = next(employee for employee in employees if employee.name == 'Q')
        assert q.ward == 3
        assert q.is_ot is True
        assert schedule.cell(row=q.row, column=1).value == 'Q（作業療法士）'
        assert schedule.cell(row=q.row, column=2).value == '勤務希望'
        assert schedule.cell(row=q.row + 1, column=2).value == '理由欄'
        assert schedule.cell(row=q.row, column=35).value == '=AS8'
        assert schedule.cell(row=q.row, column=36).value == (
            f'=COUNTIF(C{q.row + 1}:AG{q.row + 1},"年休")'
            f'+COUNTIF(C{q.row + 1}:AG{q.row + 1},"/年")/2'
            f'+COUNTIF(C{q.row + 1}:AG{q.row + 1},"年/")/2'
        )
        assert (
            schedule.row_dimensions[q.row].height
            == schedule.row_dimensions[q.row - 2].height
        )
        assert any(
            merged.min_row == q.row and merged.max_row == q.row + 1
            and merged.min_col == 1 and merged.max_col == 1
            for merged in schedule.merged_cells.ranges
        )
        p_rows = [
            row for row in range(1, schedule.max_row + 1)
            if employee_name(schedule.cell(row, 1).value) == 'P'
        ]
        assert len(p_rows) == 1
        assert all(employee.name != 'P' for employee in employees)
        assert [
            [schedule.cell(row, column).value
             for column in range(1, len(original_p_pair[0]) + 1)]
            for row in (p_rows[0], p_rows[0] + 1)
        ] == original_p_pair
        assert_active_holiday_summary_references(
            schedule, employees, day_columns(schedule),
        )
        formulas = capture_formulas(output_book)
        assert formulas
        assert all('#REF!' not in formula for formula in formulas.values())
    finally:
        output_book.close()
