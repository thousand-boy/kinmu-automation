from pathlib import Path
from unittest.mock import Mock

import pytest
from openpyxl import Workbook, load_workbook

from kinmu import application
from kinmu.models import Employee, HolidayTarget
from kinmu.validation import capture_formulas


@pytest.fixture
def save_arguments():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Sheet1'
    sheet['A4'] = 'A'
    sheet['B4'] = '勤務希望'
    sheet['B5'] = '理由欄'
    sheet['H4'] = '=1+1'
    master = workbook.create_sheet('職員マスタ')
    master.append(['職員ID', '表示名', '病棟', '専従', '専任', '作業療法士', '在籍'])
    master.append(['A', 'A', 1, '○', None, None, '○'])
    employee = Employee('A', 'A', 4, 1, True, False, False)
    yield (
        workbook, capture_formulas(workbook), {},
        [employee], [3], {'A': HolidayTarget(0, 0)},
    )
    workbook.close()


@pytest.mark.parametrize(
    'failure_stage', ['write', 'read', 'validate', 'replace'],
)
def test_failed_save_preserves_previous_output_and_removes_temporary_file(
    tmp_path, monkeypatch, save_arguments, failure_stage,
):
    output_file = tmp_path / 'result.xlsx'
    output_file.write_bytes(b'previous output')
    unrelated_temp = tmp_path / '.another-run.xlsx'
    unrelated_temp.write_bytes(b'another process owns this file')
    original_files = {
        path.name: path.read_bytes() for path in tmp_path.iterdir()
    }

    if failure_stage == 'write':
        monkeypatch.setattr(
            save_arguments[0], 'save', Mock(side_effect=OSError('disk full')),
        )
        error = OSError
    elif failure_stage == 'read':
        monkeypatch.setattr(
            application, 'load_workbook',
            Mock(side_effect=ValueError('cannot reopen workbook')),
        )
        error = ValueError
    elif failure_stage == 'validate':
        # 実際の再読み込み後に数式の不一致を検出させる。
        save_arguments[1]['Sheet1', 'H4'] = '=999'
        error = RuntimeError
    else:
        monkeypatch.setattr(
            Path, 'replace', Mock(side_effect=PermissionError('file locked')),
        )
        error = PermissionError

    with pytest.raises(error):
        application.save_safely(*save_arguments, output_file)

    assert {
        path.name: path.read_bytes() for path in tmp_path.iterdir()
    } == original_files


def test_successful_save_replaces_output_after_validation(
    tmp_path, save_arguments,
):
    output_file = tmp_path / 'result.xlsx'
    output_file.write_bytes(b'previous output')
    application.save_safely(*save_arguments, output_file)
    assert list(tmp_path.iterdir()) == [output_file]
    workbook = load_workbook(output_file)
    try:
        assert workbook['Sheet1']['H4'].value == '=1+1'
        assert workbook.calculation.fullCalcOnLoad is True
        assert workbook.calculation.forceFullCalc is True
    finally:
        workbook.close()


def test_run_rejects_input_errors_before_solving_or_saving(
    tmp_path, monkeypatch,
):
    source = (
        Path(__file__).resolve().parents[1] / 'sample' / 'kinmu_sample.xlsx'
    )
    input_file = tmp_path / 'input.xlsx'
    output_file = tmp_path / 'output.xlsx'
    workbook = load_workbook(source)
    try:
        workbook['Sheet1']['C4'] = '不明な勤務'
        workbook.save(input_file)
    finally:
        workbook.close()
    original_input = input_file.read_bytes()
    output_file.write_bytes(b'previous output')
    solve = Mock()
    save = Mock()
    monkeypatch.setattr(application, 'solve_schedule', solve)
    monkeypatch.setattr(application, 'save_safely', save)

    with pytest.raises(ValueError, match='入力エラー'):
        application.run(input_file, output_file)

    solve.assert_not_called()
    save.assert_not_called()
    assert input_file.read_bytes() == original_input
    assert output_file.read_bytes() == b'previous output'


def test_run_closes_first_workbook_when_loading_cache_fails(
    tmp_path, monkeypatch,
):
    input_file = tmp_path / 'input.xlsx'
    input_file.touch()
    workbook = Mock()
    loader = Mock(side_effect=[workbook, OSError('cache read failed')])
    monkeypatch.setattr(application, 'load_workbook', loader)

    with pytest.raises(OSError, match='cache read failed'):
        application.run(input_file, tmp_path / 'output.xlsx')

    workbook.close.assert_called_once_with()
    assert loader.call_count == 2
    assert not (tmp_path / 'output.xlsx').exists()


def test_run_does_not_overwrite_its_input(tmp_path):
    input_file = tmp_path / 'input.xlsx'
    input_file.write_bytes(b'original input')
    with pytest.raises(ValueError, match='別のパス'):
        application.run(input_file, input_file)
    assert input_file.read_bytes() == b'original input'


def test_run_rejects_output_symlink_to_input(tmp_path):
    input_file = tmp_path / 'input.xlsx'
    input_file.write_bytes(b'original input')
    output_file = tmp_path / 'alias.xlsx'
    try:
        output_file.symlink_to(input_file)
    except OSError:
        pytest.skip('シンボリックリンクを作成できない環境です。')

    with pytest.raises(ValueError, match='別のパス'):
        application.run(input_file, output_file)
    assert input_file.read_bytes() == b'original input'
    assert output_file.is_symlink()


def test_run_reports_missing_input_without_creating_output(tmp_path):
    with pytest.raises(FileNotFoundError, match='勤務表が見つかりません'):
        application.run(tmp_path / 'missing.xlsx', tmp_path / 'output.xlsx')
    assert not list(tmp_path.iterdir())


def test_main_passes_explicit_file_arguments(monkeypatch):
    run = Mock()
    monkeypatch.setattr(application, 'run', run)
    application.main(['--input', 'source.xlsx', '--output', 'result.xlsx'])
    run.assert_called_once_with(Path('source.xlsx'), Path('result.xlsx'))
